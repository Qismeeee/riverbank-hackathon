import io
import json
import logging
import zipfile
from dataclasses import replace
from datetime import date, datetime, timedelta
from pathlib import Path
from threading import Barrier, Event, Lock
from types import SimpleNamespace

import pytest

import config
import llm
from llm import LlmClient, LlmError, LlmRequest, LlmResponse
from prompts import default_prompts
from services.statements import factory
from services.statements.balance import check_balances, check_totals
from services.statements.classifiers import (
    GatedPayoutClassifier,
    JevPayoutClassifier,
    PatternPayoutClassifier,
)
from services.statements.documents import (
    CSV,
    HEIC,
    JPEG,
    PDF,
    PNG,
    WEBP,
    XLS,
    XLSX,
    StatementDocument,
    StatementLine,
    StatementSummary,
    fold,
    parse_date,
    parse_vnd,
)
from services.statements.images import ImageConfig, ImageNormalizer
from services.statements.parsers import (
    ColumnMap,
    CsvStatementParser,
    ExcelStatementParser,
    LegacyExcelStatementParser,
    PdfLayoutStatementParser,
    PdfPageConfig,
    PdfTextStatementParser,
)
from services.statements.reader import StatementConfig, StatementReader
from services.statements.storage import LocalDocumentStore, NoDocumentStore, S3DocumentStore
from services.statements.vision import VisionConfig, VisionStatementParser
from services.underwriting.models import WorkflowError
from tests.test_underwriting import (  # noqa: F401
    BANK,
    client,
    create,
    seller_headers,
    service,
)

FIXTURES = Path(__file__).parent / "fixtures" / "statements"


def image_bytes(fmt: str = "PNG", size=(40, 20), **save) -> bytes:
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", size, "white").save(out, format=fmt, **save)
    return out.getvalue()


PNG_BYTES = image_bytes()
HEADER = ["Ngày giao dịch", "Nội dung", "Ghi nợ", "Ghi có (VND)"]


def pdf_bytes(*pages: list[list[str]]) -> bytes:
    xs, height, top = [40, 130, 390, 470, 570, 610], 20, 760
    streams = []
    for rows in pages:
        ops = []
        for r, row in enumerate(rows):
            y = top - (r + 1) * height
            for c, cell in enumerate(row):
                ops.append(f"{xs[c]} {y} {xs[c + 1] - xs[c]} {height} re S")
                if cell:
                    ops.append(f"BT /F1 8 Tf {xs[c] + 3} {y + 6} Td ({cell}) Tj ET")
        streams.append("\n".join(ops).encode("latin-1"))
    count = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(count)).encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, count),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, stream in enumerate(streams):
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % (5 + 2 * i)
        )
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


def csv_bytes(rows: list[list[str]], delimiter: str = ";") -> bytes:
    return "\n".join(delimiter.join(row) for row in [HEADER, *rows]).encode("utf-8-sig")


def xlsx_bytes(*sheets: list[list]) -> bytes:
    import openpyxl

    book = openpyxl.Workbook()
    book.remove(book.active)
    for number, rows in enumerate(sheets):
        sheet = book.create_sheet(f"Sheet{number + 1}")
        for row in rows:
            sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def zip_bytes(name: str) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr(name, "x")
    return out.getvalue()


class Captured(logging.Handler):
    def __init__(self):
        super().__init__()
        self.lines: list[str] = []

    def emit(self, record):
        self.lines.append(record.getMessage())


class ScriptedVision(LlmClient):
    model = "vision-test"

    def __init__(self, *pages: dict, retries: dict[int, dict] | None = None):
        self.pages, self.retries, self.requests = list(pages), retries or {}, []
        self._lock = Lock()

    def complete_json(self, request: LlmRequest) -> LlmResponse:
        with self._lock:
            self.requests.append(request)
        number = request.user["page"]
        script = self.retries if "retry_reason" in request.user else dict(enumerate(self.pages, 1))
        return LlmResponse(script[number], "vision-test-2026")


def ocr_line(day, desc, credit=None, debit=None, balance=None) -> dict:
    return {
        "date": day,
        "raw_desc": desc,
        "credit_amount_vnd": credit,
        "debit_amount_vnd": debit,
        "balance_vnd": balance,
    }


def ocr_page(*lines) -> dict:
    return {"unreadable": False, "lines": list(lines)}


def reader(*parsers, pattern=r"\bshopee\b") -> StatementReader:
    defaults = [CsvStatementParser(), PdfTextStatementParser()]
    return StatementReader(list(parsers) or defaults, PatternPayoutClassifier(pattern))


@pytest.mark.parametrize(
    "raw,value",
    [
        ("70.000.000", 70_000_000),
        ("70,000,000", 70_000_000),
        ("70.000.000,00", 70_000_000),
        ("1,500.50", 1_500),
        ("(2.000)", 2_000),
        ("", None),
        ("-", None),
        (12_000, 12_000),
    ],
)
def test_parse_vnd_reads_vietnamese_and_english_separators(raw, value):
    assert parse_vnd(raw) == value


def test_parse_date_and_fold():
    assert parse_date("05/10/2026 08:12") == "2026-10-05"
    assert parse_date("2026-10-05") == "2026-10-05"
    assert parse_date("31/02/2026") is None
    assert fold("  Ngày GIAO   dịch Đà ") == "ngay giao dich da"


def test_sniff_uses_magic_bytes_not_client_type():
    assert StatementDocument.sniff(b"%PDF-1.4 ...").media_type == PDF
    assert StatementDocument.sniff(PNG_BYTES).media_type == PNG
    assert StatementDocument.sniff("a;b\n".encode()).media_type == CSV
    with pytest.raises(WorkflowError) as exc:
        StatementDocument.sniff(b"\xff\xfe\x00\x81binary")
    assert (exc.value.code, exc.value.status) == ("STATEMENT_FORMAT_UNSUPPORTED", 415)


def test_csv_parser_maps_headers_and_skips_balance_rows():
    content = csv_bytes(
        [
            ["05/10/2026", "SHOPEE VN THANH TOAN DOANH THU", "", "70.000.000"],
            ["06/10/2026", "Chuyen tien mua hang", "5.000.000", ""],
            ["", "Số dư cuối kỳ", "", ""],
        ]
    )
    parsed = CsvStatementParser().parse(StatementDocument.sniff(content))
    assert [(x.date, x.credit, x.debit) for x in parsed.lines] == [
        ("2026-10-05", 70_000_000, None),
        ("2026-10-06", None, 5_000_000),
    ]


def test_csv_parser_reads_signed_amount_column():
    text = 'Date,Description,Amount\n2026-10-05,Shopee payout,"1,200,000"\n2026-10-06,Fee,-50000\n'
    lines = CsvStatementParser().parse(StatementDocument.sniff(text.encode())).lines
    assert [(x.credit, x.debit) for x in lines] == [(1_200_000, None), (None, 50_000)]


def test_sniff_accepts_workbook_and_rejects_other_zips():
    assert StatementDocument.sniff(xlsx_bytes([["a"]])).media_type == XLSX
    for content in (zip_bytes("word/document.xml"), b"PK\x03\x04broken"):
        with pytest.raises(WorkflowError) as exc:
            StatementDocument.sniff(content)
        assert (exc.value.code, exc.value.status) == ("STATEMENT_FORMAT_UNSUPPORTED", 415)


def test_excel_parser_reads_native_cells_across_sheets():
    content = xlsx_bytes(
        [["Sao kê tài khoản"], HEADER, [datetime(2026, 10, 5, 8, 12), "SHOPEE VN", None, 70e6]],
        [
            ["Transaction date", "Description", "Amount"],
            [date(2026, 10, 6), "Fee", -50_000],
            ["07/10/2026", "Shopee payout", "1.200.000"],
            [None, "Closing balance", None],
        ],
    )
    parsed = ExcelStatementParser().parse(StatementDocument.sniff(content))
    assert (parsed.parser, parsed.pages) == ("xlsx", 2)
    assert [(x.date, x.raw_desc, x.credit, x.debit) for x in parsed.lines] == [
        ("2026-10-05", "SHOPEE VN", 70_000_000, None),
        ("2026-10-06", "Fee", None, 50_000),
        ("2026-10-07", "Shopee payout", 1_200_000, None),
    ]


def test_excel_parser_reports_missing_header_and_broken_workbook():
    empty = ExcelStatementParser().parse(StatementDocument.sniff(xlsx_bytes([["a", "b"]])))
    assert (empty.lines, empty.notes) == ([], ["NO_STATEMENT_HEADER"])
    broken = ExcelStatementParser().parse(StatementDocument(zip_bytes("xl/workbook.xml"), XLSX))
    assert (broken.lines, broken.notes) == ([], ["WORKBOOK_UNREADABLE"])


def test_excel_parser_stops_at_row_limit():
    rows = [HEADER, *[["05/10/2026", "SHOPEE", "", "1.000"]] * 5]
    parsed = ExcelStatementParser(max_rows=3).parse(StatementDocument.sniff(xlsx_bytes(rows)))
    assert len(parsed.lines) == 2


def test_pdf_text_layer_parses_table_without_model():
    content = pdf_bytes(
        [
            ["Ngay giao dich", "Noi dung", "Ghi no", "Ghi co"],
            ["05/10/2026", "SHOPEE VN THANH TOAN", "", "70.000.000"],
            ["06/10/2026", "Phi dich vu", "11.000", ""],
        ]
    )
    parsed = PdfTextStatementParser().parse(StatementDocument.sniff(content))
    assert parsed.parser == "pdf_text" and parsed.pages == 1
    assert [(x.date, x.raw_desc, x.credit, x.debit) for x in parsed.lines] == [
        ("2026-10-05", "SHOPEE VN THANH TOAN", 70_000_000, None),
        ("2026-10-06", "Phi dich vu", None, 11_000),
    ]


def test_scanned_pdf_falls_back_to_vision_page_by_page():
    vision_client = ScriptedVision(
        {"unreadable": False, "lines": [ocr_line("2026-10-05", "SHOPEE VN", 70_000_000)]}
    )
    vision = VisionStatementParser(vision_client, default_prompts().get("statement-ocr"))
    scanned = pdf_bytes([["", "", "", ""]])
    read = reader(CsvStatementParser(), PdfTextStatementParser(), vision).read(scanned)
    assert read.parsed.parser == "vision" and read.parsed.model == "vision-test-2026"
    assert read.parsed.prompt_version.startswith("statement-ocr@")
    request = vision_client.requests[0]
    assert request.images[0].startswith("data:image/jpeg;base64,")
    assert request.user == {"page": 1, "pages": 1} and request.temperature == 0


def test_unreadable_document_lists_every_attempt():
    with pytest.raises(WorkflowError) as exc:
        reader().read(pdf_bytes([["", "", "", ""]]))
    assert exc.value.code == "STATEMENT_UNREADABLE"
    assert exc.value.details == [{"parser": "pdf_text", "notes": ["NO_TEXT_LAYER"]}]
    with pytest.raises(WorkflowError) as exc:
        reader().read(PNG_BYTES)
    assert (exc.value.code, exc.value.status) == ("STATEMENT_FORMAT_UNSUPPORTED", 415)


def test_vision_rejects_too_many_pages_and_maps_llm_errors():
    class Down(LlmClient):
        model = "down"

        def complete_json(self, request):
            raise LlmError("LLM_UNAVAILABLE")

    prompt = default_prompts().get("statement-ocr")
    with pytest.raises(WorkflowError) as exc:
        VisionStatementParser(Down(), prompt).parse(StatementDocument(PNG_BYTES, PNG))
    assert (exc.value.code, exc.value.status, exc.value.details) == (
        "STATEMENT_OCR_UNAVAILABLE",
        503,
        ["LLM_UNAVAILABLE"],
    )
    parser = VisionStatementParser(Down(), prompt, VisionConfig(max_pages=0))
    with pytest.raises(WorkflowError, match="STATEMENT_TOO_MANY_PAGES"):
        parser.parse(StatementDocument.sniff(pdf_bytes([["", "", "", ""]])))


def test_reader_limits_size_and_empty_files():
    small = StatementReader([CsvStatementParser()], PatternPayoutClassifier("x"))
    small.cfg = StatementConfig(max_bytes=4)
    with pytest.raises(WorkflowError) as exc:
        small.read(b"12345")
    assert (exc.value.code, exc.value.status) == ("STATEMENT_TOO_LARGE", 413)
    with pytest.raises(WorkflowError, match="STATEMENT_EMPTY"):
        small.read(b"")


def test_reconcile_sums_payout_credits_in_period_and_redacts_pii():
    content = csv_bytes(
        [
            ["05/10/2026", "SHOPEE VN THANH TOAN", "", "70.000.000"],
            ["12/10/2026", "Shopee payout tuan 41", "", "69.000.000"],
            ["13/10/2026", "Nguyen Van A chuyen tien 0912345678", "", "50.000.000"],
            ["14/10/2026", "SHOPEE VN hoan phi", "300.000", ""],
            ["20/10/2026", "SHOPEE VN THANH TOAN", "", "71.000.000"],
        ]
    )
    rb = reader()
    read = rb.read(content)
    result = rb.reconcile(read, date(2026, 10, 5), date(2026, 10, 19))
    assert result["marketplace_total"] == 139_000_000
    assert [r["is_payout"] for r in result["statement_lines"]] == [True, True, False, False, True]
    assert [r["in_period"] for r in result["statement_lines"]] == [True, True, True, True, False]
    assert "[REDACTED_PHONE]" in result["statement_lines"][2]["raw_desc"]
    assert result["document_ref"] == read.document.ref
    assert result["coverage"] == {"first_date": "2026-10-05", "last_date": "2026-10-20"}
    assert result["classifier"]["name"] == "pattern"


def test_reconcile_refuses_payout_lines_without_date_or_amount():
    vision_client = ScriptedVision(
        {
            "unreadable": False,
            "lines": [
                ocr_line("2026-10-05", "SHOPEE VN", 70_000_000),
                ocr_line(None, "SHOPEE VN", 1_000),
                ocr_line("2026-10-06", "SHOPEE VN", None),
            ],
        }
    )
    rb = reader(VisionStatementParser(vision_client, default_prompts().get("statement-ocr")))
    with pytest.raises(WorkflowError) as exc:
        rb.reconcile(rb.read(PNG_BYTES), date(2026, 10, 5), date(2026, 10, 19))
    assert (exc.value.code, exc.value.details) == ("STATEMENT_LINES_UNREADABLE", [1, 2])


def test_jev_classifier_batches_and_fails_closed():
    calls = []

    def decide(state, questions, timeout):
        calls.append(questions)
        return {
            name: {"noul": 0.9 if "SHOPEE" in q["instructions"] else 0.1}
            for name, q in questions.items()
        }

    labels = JevPayoutClassifier(batch_size=2, decide=decide).classify(
        ["SHOPEE VN", "ban be", "SHOPEE payout"]
    )
    assert [(x.is_payout, x.probability) for x in labels] == [
        (True, 0.9),
        (False, 0.1),
        (True, 0.9),
    ]
    assert [len(q) for q in calls] == [2, 1]

    def down(*args, **kwargs):
        raise TimeoutError

    with pytest.raises(WorkflowError) as exc:
        JevPayoutClassifier(decide=down).classify(["x"])
    assert (exc.value.code, exc.value.status) == ("STATEMENT_CLASSIFIER_UNAVAILABLE", 503)


def test_factory_validates_config(monkeypatch):
    assert [p.name for p in factory.build_parsers("off")] == [
        "csv",
        "xlsx",
        "xls",
        "pdf_text",
        "pdf_layout",
    ]
    with pytest.raises(WorkflowError, match="STATEMENT_OCR_UNKNOWN"):
        factory.build_parsers("tesseract")
    with pytest.raises(WorkflowError, match="STATEMENT_CLASSIFIER_UNKNOWN"):
        factory.build_classifier("magic")
    monkeypatch.setattr(config, "OPENAI_API_KEY", "")
    monkeypatch.setattr(config, "STATEMENT_OCR_PROVIDER", "openai")
    with pytest.raises(WorkflowError, match="STATEMENT_OCR_NOT_CONFIGURED"):
        factory.build_parsers("llm")
    monkeypatch.setattr(config, "OPENROUTER_API_KEY", "or-test")
    assert factory.build_classifier("jev").name == "pattern+jev"
    monkeypatch.setattr(config, "OPENROUTER_API_KEY", "")
    with pytest.raises(WorkflowError, match="OPENROUTER_KEY_REQUIRED"):
        factory.build_classifier("jev")


def test_llm_client_sends_images_as_content_parts():
    llm_client = llm.OpenAICompatibleClient("sk-test", "m")
    calls = []
    reply = SimpleNamespace(
        choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps({"lines": []})))],
        model="m",
        usage=None,
    )
    llm_client._client = SimpleNamespace(
        chat=SimpleNamespace(
            completions=SimpleNamespace(create=lambda **kw: calls.append(kw) or reply)
        )
    )
    llm_client.complete_json(
        LlmRequest("s", {"page": 1}, {}, "n", images=("data:image/png;base64,AA",))
    )
    content = calls[0]["messages"][1]["content"]
    assert content[0] == {"type": "text", "text": '{"page": 1}'}
    assert content[1]["image_url"]["url"] == "data:image/png;base64,AA"


def bad_request() -> Exception:
    import httpx
    from openai import BadRequestError

    response = httpx.Response(400, request=httpx.Request("POST", "https://llm.test"))
    return BadRequestError("response_format not supported", response=response, body=None)


def scripted_openai(*replies):
    calls, queue = [], list(replies)

    def create(**kw):
        calls.append(kw)
        item = queue.pop(0)
        if isinstance(item, Exception):
            raise item
        return SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content=item))], model="m", usage=None
        )

    return calls, SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))


def test_llm_client_falls_back_to_json_mode_when_schema_is_rejected():
    llm_client = llm.OpenAICompatibleClient("sk-test", "m")
    calls, llm_client._client = scripted_openai(bad_request(), '```json\n{"lines": []}\n```')
    schema = {"type": "object", "properties": {"lines": {"type": "array"}}}
    response = llm_client.complete_json(LlmRequest("sys", {"page": 1}, schema, "n"))
    assert response.content == {"lines": []}
    assert calls[0]["response_format"]["type"] == "json_schema"
    assert calls[1]["response_format"] == {"type": "json_object"}
    assert calls[1]["messages"][0]["content"].startswith("sys\n\n" + llm.SCHEMA_HINT)
    _, llm_client._client = scripted_openai(bad_request(), bad_request())
    with pytest.raises(LlmError, match="LLM_UNAVAILABLE"):
        llm_client.complete_json(LlmRequest("sys", {}, schema, "n"))


@pytest.fixture
def upload_service(service):  # noqa: F811
    service.statements = reader()
    return service


def period_rows(case) -> tuple[str, str, list[str]]:
    weeks = case.runs[-1]["input_snapshot"]["features"]["weekly_payouts"]
    start = date.fromisoformat(weeks[-2]["week_start"])
    end = date.fromisoformat(weeks[-1]["week_start"]) + timedelta(days=7)
    inside = [(start + timedelta(days=d)).strftime("%d/%m/%Y") for d in (1, 8)]
    return start.isoformat(), end.isoformat(), inside


def upload(api, case, content, headers=BANK, version=None, name="saoke.csv"):
    start, end, _ = period_rows(case)
    return api.post(
        f"/v1/cases/{case.case_id}/statements/upload",
        headers={k: v for k, v in headers.items() if k != "Idempotency-Key"},
        data={
            "expected_version": str(version or case.version),
            "period_start": start,
            "period_end": end,
        },
        files={"file": (name, content, "application/octet-stream")},
    )


@pytest.mark.parametrize(
    "second,status", [("49.000.000", "PENDING_REVIEW"), ("48.999.999", "ESCALATED")]
)
def test_uploaded_statement_drives_r05_through_existing_threshold(
    client,  # noqa: F811
    upload_service,
    second,
    status,
):
    case = upload_service.repository.get(create(client, "ducthinh.tech")["case_id"])
    _, _, (first_day, second_day) = period_rows(case)
    content = csv_bytes(
        [
            [first_day, "SHOPEE VN THANH TOAN", "", "70.000.000"],
            [second_day, "SHOPEE VN THANH TOAN", "", second],
            [second_day, "Ban be chuyen khoan", "", "40.000.000"],
        ]
    )
    response = upload(client, case, content, seller_headers("ducthinh.tech"))
    assert response.status_code == 200, response.text
    updated = upload_service.repository.get(case.case_id)
    assert updated.status == status and len(updated.runs) == 2
    statement = updated.runs[-1]["statement"]
    assert statement["source"] == "uploaded_file" and statement["submitted_by"] == "ducthinh.tech"
    assert statement["platform_total"] == 140_000_000
    assert statement["marketplace_total"] == 70_000_000 + parse_vnd(second)
    assert statement["extraction"]["parser"] == "csv"
    assert any(e["code"] == "E8" for e in updated.evidence)
    assert updated.events[-1]["event_type"] == "STATEMENT_RECORDED"


def test_uploaded_workbook_reconciles_like_csv(client, upload_service):  # noqa: F811
    upload_service.statements = reader(CsvStatementParser(), ExcelStatementParser())
    case = upload_service.repository.get(create(client, "ducthinh.tech")["case_id"])
    _, _, (first_day, second_day) = period_rows(case)
    content = xlsx_bytes(
        [
            HEADER,
            [first_day, "SHOPEE VN THANH TOAN", None, 70_000_000],
            [second_day, "SHOPEE VN THANH TOAN", None, 70_000_000],
        ]
    )
    response = upload(client, case, content, name="saoke.xlsx")
    assert response.status_code == 200, response.text
    statement = upload_service.repository.get(case.case_id).runs[-1]["statement"]
    assert statement["extraction"]["parser"] == "xlsx"
    assert statement["marketplace_total"] == 140_000_000


def test_upload_scope_version_and_state_checks(client, upload_service):  # noqa: F811
    case = upload_service.repository.get(create(client, "ducthinh.tech")["case_id"])
    content = csv_bytes([["05/10/2026", "SHOPEE VN", "", "1.000"]])
    assert upload(client, case, content, seller_headers("minhchau.shop")).status_code == 403
    assert upload(client, case, content, version=case.version + 1).status_code == 409
    assert upload(client, case, b"\x00\x81\xff").status_code == 415
    ready = upload_service.repository.get(create(client, "minhchau.shop")["case_id"])
    response = upload(client, ready, content)
    assert response.json()["detail"]["code"] == "STATEMENT_NOT_REQUESTED"
    assert upload_service.repository.get(case.case_id).version == case.version


def test_upload_without_reader_is_unavailable(client, service):  # noqa: F811
    case = service.repository.get(create(client, "ducthinh.tech")["case_id"])
    response = upload(client, case, csv_bytes([]))
    assert response.status_code == 503
    assert response.json()["detail"]["code"] == "STATEMENT_READER_UNAVAILABLE"


@pytest.mark.parametrize(
    "header,columns",
    [
        (
            [
                "Ngày GD/Trans Date",
                "Nội dung/Description",
                "Số tiền ghi nợ/Debit",
                "Số tiền ghi có/Credit",
                "Số dư/Balance",
            ],
            {"date": 0, "desc": 1, "debit": 2, "credit": 3},
        ),
        (
            ["Ngày hiệu lực", "Diễn giải", "Phát sinh nợ", "Phát sinh có"],
            {"date": 0, "desc": 1, "debit": 2, "credit": 3},
        ),
        (
            ["STT", "Thời gian giao dịch", "Nội dung chuyển khoản", "Số tiền"],
            {"date": 1, "desc": 2, "amount": 3},
        ),
        (
            ["No.", "Date", "Description", "Debit", "Credit", "Balance"],
            {"date": 1, "desc": 2, "debit": 3, "credit": 4},
        ),
        (["STT", "Mô tả", "Nợ", "Có"], {"desc": 1, "debit": 2, "credit": 3}),
        (["Số dư", "Ghi chú", "Mã GD"], None),
    ],
)
def test_header_detection_covers_bank_variants(header, columns):
    found = ColumnMap.detect(header)
    assert (found.columns if found else None) == columns


def test_unlabelled_date_column_is_inferred_from_content():
    text = "STT;;Mô tả;Có\n1;05/10/2026;SHOPEE VN;70.000.000\n2;06/10/2026;Ban be;1.000\n"
    lines = CsvStatementParser().parse(StatementDocument.sniff(text.encode())).lines
    assert [(x.date, x.credit) for x in lines] == [
        ("2026-10-05", 70_000_000),
        ("2026-10-06", 1_000),
    ]


def jev_rejecting(*names: str):
    calls = []

    def decide(state, questions, timeout):
        calls.append(questions)
        return {
            key: {"noul": 0.05 if any(n in q["instructions"] for n in names) else 0.95}
            for key, q in questions.items()
        }

    return decide, calls


def test_gated_classifier_only_asks_jev_about_platform_candidates():
    decide, calls = jev_rejecting("NGUYEN VAN A")
    gated = GatedPayoutClassifier(
        PatternPayoutClassifier("shopee"), JevPayoutClassifier(decide=decide)
    )
    labels = gated.classify(
        ["SHOPEE VN quyet toan", "NGUYEN VAN A chuyen tien hang shopee", "Luong thang 9"]
    )
    assert [x.is_payout for x in labels] == [True, False, False]
    assert [len(q) for q in calls] == [2]


@pytest.mark.parametrize(
    "fixture,parser,pages,total",
    [
        ("bilingual_grid.pdf", "pdf_text", 2, 12 * (70_000_000 + 69_500_000)),
        ("unruled_text.pdf", "pdf_layout", 1, 70_000_000 + 69_500_000),
    ],
)
def test_vietnamese_pdf_fixtures_reconcile_without_model(fixture, parser, pages, total):
    decide, _ = jev_rejecting("NGUYEN VAN A")
    rb = StatementReader(
        [CsvStatementParser(), PdfTextStatementParser(), PdfLayoutStatementParser()],
        GatedPayoutClassifier(
            PatternPayoutClassifier("shopee"), JevPayoutClassifier(decide=decide)
        ),
    )
    read = rb.read((FIXTURES / fixture).read_bytes())
    result = rb.reconcile(read, date(2026, 10, 5), date(2026, 10, 19))
    assert (read.parsed.parser, read.parsed.pages) == (parser, pages)
    assert result["marketplace_total"] == total
    first = result["statement_lines"][0]
    assert first["raw_desc"].endswith("tuan 40 shop ABC") and first["debit"] is None


def test_layout_parser_reads_grid_pdf_across_pages_too():
    parsed = PdfLayoutStatementParser().parse(
        StatementDocument.sniff((FIXTURES / "bilingual_grid.pdf").read_bytes())
    )
    payouts = [x for x in parsed.lines if x.credit and x.date]
    assert len(payouts) == 36 and parsed.pages == 2


def test_statement_threshold_comes_from_policy_config(client, upload_service):  # noqa: F811
    upload_service.policy.cfg = replace(upload_service.policy.cfg, statement_mismatch_bps=0)
    case = upload_service.repository.get(create(client, "ducthinh.tech")["case_id"])
    _, _, (first_day, second_day) = period_rows(case)
    content = csv_bytes(
        [
            [first_day, "SHOPEE VN THANH TOAN", "", "70.000.000"],
            [second_day, "SHOPEE VN THANH TOAN", "", "69.999.999"],
        ]
    )
    assert upload(client, case, content).status_code == 200
    statement = upload_service.repository.get(case.case_id).runs[-1]["statement"]
    assert (statement["threshold_bps"], statement["matched"]) == (0, False)


def line(credit=None, debit=None, balance=None, page=None) -> StatementLine:
    return StatementLine("2026-10-05", "x", credit, debit, balance, page)


def test_balance_check_accepts_both_orders_and_flags_bad_digits():
    oldest = [line(balance=100), line(credit=50, balance=150), line(debit=30, balance=120)]
    assert check_balances(oldest).payload() == {
        "checked_pairs": 2,
        "mismatched_lines": [],
        "order": "oldest_first",
        "totals": None,
    }
    assert check_balances(oldest[::-1]).order == "newest_first"
    misread = [line(balance=100), line(credit=80, balance=150), line(debit=30, balance=120)]
    assert check_balances(misread).mismatched_lines == (1,)
    gap = [line(balance=100, page=1), line(credit=5, balance=999, page=3)]
    assert check_balances(gap).checked_pairs == 0
    assert check_balances([line(credit=1), line(credit=2)]).order is None


def test_text_parsers_read_signed_balance_column():
    header = "Ngày;Nội dung;Ghi nợ;Ghi có;Số dư\n"
    rows = "05/10/2026;SHOPEE;;70.000.000;-1.000.000\n06/10/2026;Phi;11.000;;-1.011.000\n"
    lines = CsvStatementParser().parse(StatementDocument.sniff((header + rows).encode())).lines
    assert [x.balance for x in lines] == [-1_000_000, -1_011_000]
    assert check_balances(lines).mismatches == 0


def test_reconcile_records_balance_check():
    content = "Ngày;Nội dung;Ghi nợ;Ghi có;Số dư\n05/10/2026;SHOPEE VN;;70.000.000;70.000.000\n"
    rb = reader()
    result = rb.reconcile(rb.read(content.encode()), date(2026, 10, 5), date(2026, 10, 12))
    assert result["extraction"]["balance_check"]["checked_pairs"] == 0
    assert result["extraction"]["image_pages"] == []


def test_vision_rereads_page_sharper_when_balance_breaks():
    wrong = ocr_page(
        ocr_line("2026-10-05", "SHOPEE VN", 70_000_000, balance=70_000_000),
        ocr_line("2026-10-06", "SHOPEE VN", 10_000_000, balance=89_000_000),
    )
    right = ocr_page(
        ocr_line("2026-10-05", "SHOPEE VN", 70_000_000, balance=70_000_000),
        ocr_line("2026-10-06", "SHOPEE VN", 19_000_000, balance=89_000_000),
    )
    ocr = ScriptedVision(wrong, retries={1: right})
    vision = VisionStatementParser(ocr, default_prompts().get("statement-ocr"))
    parsed = vision.parse(StatementDocument.sniff(pdf_bytes([["", "", "", ""]])))
    assert [x.credit for x in parsed.lines] == [70_000_000, 19_000_000]
    assert parsed.notes == ["PAGE_1_RETRIED"] and parsed.lines[0].page == 1
    first, second = ocr.requests
    assert second.user == {"page": 1, "pages": 1, "retry_reason": "BALANCE_MISMATCH"}
    assert len(second.images[0]) > len(first.images[0])


def test_vision_keeps_first_read_when_retry_is_not_better():
    wrong = ocr_page(
        ocr_line("2026-10-05", "A", 1, balance=1), ocr_line("2026-10-06", "B", 1, balance=9)
    )
    ocr = ScriptedVision(wrong, retries={1: ocr_page()})
    parsed = VisionStatementParser(ocr, default_prompts().get("statement-ocr")).parse(
        StatementDocument(PNG_BYTES, PNG)
    )
    assert len(parsed.lines) == 2 and parsed.notes == ["PAGE_1_RETRIED"]


def test_mixed_pdf_sends_only_image_pages_to_vision():
    text_page = [
        ["Ngay giao dich", "Noi dung", "Ghi no", "Ghi co", "So du"],
        ["05/10/2026", "SHOPEE VN THANH TOAN", "", "70.000", "70.000"],
    ]
    content = pdf_bytes(text_page, [["", "", "", "", ""]], text_page)
    ocr = ScriptedVision(None, ocr_page(ocr_line("2026-10-08", "SHOPEE", 5_000, balance=75_000)))
    vision = VisionStatementParser(ocr, default_prompts().get("statement-ocr"))
    pages = PdfPageConfig(defer_images=True)
    read = reader(PdfTextStatementParser(pages), vision).read(content)
    assert read.parsed.parser == "pdf_text+vision" and read.parsed.image_pages == (2,)
    assert [(x.page, x.credit) for x in read.parsed.lines] == [(1, 70_000), (2, 5_000), (3, 70_000)]
    assert [r.user["page"] for r in ocr.requests] == [2]
    assert check_balances(read.parsed.lines).mismatched_lines == (2,)
    plain = reader(PdfTextStatementParser()).read(content)
    assert plain.parsed.parser == "pdf_text" and len(plain.parsed.lines) == 2


def test_mixed_pdf_without_vision_says_which_pages_were_skipped():
    content = pdf_bytes(
        [["Ngay", "Noi dung", "Ghi no", "Ghi co"], ["05/10/2026", "SHOPEE", "", "1.000"]],
        [["", "", "", ""]],
    )
    pages = PdfPageConfig(defer_images=True)
    read = reader(PdfTextStatementParser(pages)).read(content)
    assert read.parsed.notes == ["PAGE_2_NOT_READ"]


def fake_page(number, words, images=(), chars=(1,)):
    return SimpleNamespace(
        page_number=number,
        chars=list(chars),
        width=612,
        height=792,
        images=list(images),
        extract_words=lambda: [{"text": w} for w in words],
    )


def test_image_page_detection_skips_watermarked_text_pages():
    cover = {"x0": 0, "x1": 600, "top": 0, "bottom": 700}
    scanned = fake_page(1, ["NGAN", "HANG", "A"], [cover])
    watermarked = fake_page(2, ["05/10/2026", "SHOPEE", "70.000"], [cover])
    plain = fake_page(3, ["05/10/2026"])
    empty = fake_page(4, [], chars=())
    parser = PdfTextStatementParser(PdfPageConfig(defer_images=True))
    assert parser.image_pages([scanned, watermarked, plain, empty]) == (1, 4)
    assert PdfLayoutStatementParser().image_pages([scanned]) == ()


def test_vision_reads_pages_in_parallel_but_keeps_order():
    pages = [ocr_page(ocr_line(f"2026-10-0{n}", f"P{n}", n)) for n in range(1, 5)]
    ocr = ScriptedVision(*pages)
    vision = VisionStatementParser(ocr, default_prompts().get("statement-ocr"))
    content = pdf_bytes(*[[["", "", "", ""]]] * 4)
    parsed = vision.parse(StatementDocument.sniff(content))
    assert [x.raw_desc for x in parsed.lines] == ["P1", "P2", "P3", "P4"]
    assert parsed.pages == 4


def test_vision_rereads_a_page_while_other_pages_are_still_reading():
    retried = Event()

    class Slow(ScriptedVision):
        def complete_json(self, request):
            if "retry_reason" in request.user:
                retried.set()
            elif request.user["page"] == 2:
                assert retried.wait(5), "page 2 blocked the retry of page 1"
            return super().complete_json(request)

    wrong = ocr_page(
        ocr_line("2026-10-05", "A", 1, balance=1), ocr_line("2026-10-06", "B", 1, balance=9)
    )
    right = ocr_page(
        ocr_line("2026-10-05", "A", 1, balance=1), ocr_line("2026-10-06", "B", 8, balance=9)
    )
    ocr = Slow(wrong, ocr_page(ocr_line("2026-10-07", "C", 5)), retries={1: right})
    vision = VisionStatementParser(ocr, default_prompts().get("statement-ocr"))
    parsed = vision.parse(StatementDocument.sniff(pdf_bytes(*[[["", "", "", ""]]] * 2)))
    assert [x.credit for x in parsed.lines] == [1, 8, 5]
    assert parsed.notes == ["PAGE_1_RETRIED"]


def test_vision_reads_all_pages_at_once_and_logs_timings():
    gate = Barrier(8, timeout=5)

    class Together(ScriptedVision):
        def complete_json(self, request):
            gate.wait()
            return super().complete_json(request)

    pages = [ocr_page(ocr_line("2026-10-05", f"P{n}", n)) for n in range(1, 9)]
    vision = VisionStatementParser(Together(*pages), default_prompts().get("statement-ocr"))
    captured, log = Captured(), logging.getLogger("riverbank.ocr")
    log.addHandler(captured)
    try:
        parsed = vision.parse(StatementDocument.sniff(pdf_bytes(*[[["", "", "", ""]]] * 8)))
    finally:
        log.removeHandler(captured)
    assert [x.raw_desc for x in parsed.lines] == [f"P{n}" for n in range(1, 9)]
    events = [json.loads(line) for line in captured.lines]
    assert sorted(e["page"] for e in events if e["event"] == "ocr_page") == list(range(1, 9))
    document = events[-1]
    assert (document["event"], document["pages"], document["retried"]) == ("ocr_document", 8, 0)
    assert document["seconds"] >= 0 and document["model"] == "vision-test"


def heic_bytes() -> bytes:
    from pillow_heif import register_heif_opener

    register_heif_opener()
    return image_bytes("HEIF")


def rotated_jpeg() -> bytes:
    from PIL import Image

    exif = Image.Exif()
    exif[0x0112] = 6
    return image_bytes("JPEG", size=(40, 20), exif=exif)


def test_sniff_recognises_phone_and_legacy_formats():
    assert StatementDocument.sniff(image_bytes("WEBP")).media_type == WEBP
    assert StatementDocument.sniff(heic_bytes()).media_type == HEIC
    legacy = (FIXTURES / "legacy.xls").read_bytes()
    assert StatementDocument.sniff(legacy).media_type == XLS


def test_images_are_turned_upright_shrunk_and_sent_as_jpeg():
    from PIL import Image

    normalizer = ImageNormalizer(ImageConfig(max_side=30))
    media_type, upright = normalizer.normalize(rotated_jpeg())
    assert media_type == JPEG
    assert Image.open(io.BytesIO(upright)).size == (15, 30)
    _, shrunk = normalizer.normalize(image_bytes(size=(300, 60)))
    assert Image.open(io.BytesIO(shrunk)).size == (30, 6)
    _, decoded = normalizer.normalize(heic_bytes())
    assert Image.open(io.BytesIO(decoded)).format == "JPEG"
    with pytest.raises(WorkflowError) as exc:
        normalizer.normalize(b"RIFF\x00\x00\x00\x00WEBPbroken")
    assert (exc.value.code, exc.value.status) == ("STATEMENT_IMAGE_UNREADABLE", 422)


def test_vision_sends_phone_photos_as_jpeg():
    ocr = ScriptedVision(ocr_page(ocr_line("2026-10-05", "SHOPEE", 1_000)))
    vision = VisionStatementParser(ocr, default_prompts().get("statement-ocr"))
    parsed = vision.parse(StatementDocument.sniff(heic_bytes()))
    assert ocr.requests[0].images[0].startswith("data:image/jpeg;base64,")
    assert parsed.lines[0].page == 1


def test_legacy_xls_reads_dates_summary_rows_and_skips_totals():
    parsed = LegacyExcelStatementParser().parse(
        StatementDocument.sniff((FIXTURES / "legacy.xls").read_bytes())
    )
    assert parsed.parser == "xls"
    assert [(x.date, x.credit, x.debit) for x in parsed.lines] == [
        ("2026-10-05", 70_000_000, None),
        ("2026-10-06", None, 40_000_000),
        ("2026-10-12", 69_500_000, None),
        ("2026-10-13", None, 11_000),
    ]
    assert parsed.summary == StatementSummary(10_000_000, 109_489_000)
    check = check_balances(parsed.lines, parsed.summary)
    assert check.mismatches == 0 and check.totals.matched


def test_legacy_parser_reports_unreadable_ole_files():
    broken = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64
    parsed = LegacyExcelStatementParser().parse(StatementDocument.sniff(broken))
    assert parsed.notes == ["WORKBOOK_UNREADABLE"]


def test_opening_and_closing_rows_check_totals_without_balance_column():
    text = (
        "Ngày;Nội dung;Ghi nợ;Ghi có\n"
        ";Số dư đầu kỳ;;5.000.000\n"
        "05/10/2026;SHOPEE VN;;70.000.000\n"
        "06/10/2026;Phi;11.000;\n"
        ";Tổng cộng;11.000;70.000.000\n"
        ";Số dư cuối kỳ;;74.989.000\n"
    )
    parsed = CsvStatementParser().parse(StatementDocument.sniff(text.encode()))
    assert len(parsed.lines) == 2
    assert parsed.summary == StatementSummary(5_000_000, 74_989_000)
    assert check_totals(parsed.lines, parsed.summary).matched
    misread = [replace(parsed.lines[0], credit=7_000_000), parsed.lines[1]]
    assert check_totals(misread, parsed.summary).payload() == {
        "opening": 5_000_000,
        "closing": 74_989_000,
        "expected_change": 69_989_000,
        "read_change": 6_989_000,
        "matched": False,
    }
    assert check_totals(parsed.lines, StatementSummary(5_000_000)) is None


def test_vision_merges_page_summaries_and_retries_on_total_mismatch():
    first = {**ocr_page(ocr_line("2026-10-05", "A", 1_000)), "opening_balance_vnd": 500}
    wrong = {**ocr_page(ocr_line("2026-10-06", "B", 2_000)), "closing_balance_vnd": 9_999}
    right = {**ocr_page(ocr_line("2026-10-06", "B", 2_000)), "closing_balance_vnd": 3_500}
    ocr = ScriptedVision(first, wrong, retries={2: right})
    vision = VisionStatementParser(ocr, default_prompts().get("statement-ocr"))
    parsed = vision.parse(StatementDocument.sniff(pdf_bytes(*[[["", "", "", ""]]] * 2)))
    assert parsed.summary == StatementSummary(500, 9_999)
    assert parsed.notes == []
    single = {**right, "opening_balance_vnd": 500, "closing_balance_vnd": 2_500}
    bad = {**single, "closing_balance_vnd": 9_999}
    ocr = ScriptedVision(bad, retries={1: single})
    parsed = VisionStatementParser(ocr, default_prompts().get("statement-ocr")).parse(
        StatementDocument(PNG_BYTES, PNG)
    )
    assert parsed.summary == StatementSummary(500, 2_500) and parsed.notes == ["PAGE_1_RETRIED"]


def test_ocr_page_limit_comes_from_config(monkeypatch):
    monkeypatch.setattr(config, "OPENROUTER_API_KEY", "or-test")
    monkeypatch.setattr(config, "STATEMENT_OCR_PROVIDER", "openrouter")
    monkeypatch.setattr(config, "STATEMENT_OCR_MAX_PAGES", 3)
    vision = factory.build_parsers("llm")[-1]
    assert vision.cfg.max_pages == 3


class FakeS3:
    def __init__(self, error=None):
        self.error, self.calls, self.objects = error, [], {}

    def put_object(self, **kw):
        self.calls.append(kw)
        if self.error:
            raise self.error
        self.objects[(kw["Bucket"], kw["Key"])] = kw
        return {"VersionId": "v1"}

    def get_object(self, **kw):
        self.calls.append(kw)
        if self.error:
            raise self.error
        stored = self.objects[(kw["Bucket"], kw["Key"])]
        return {"Body": io.BytesIO(stored["Body"]), "ContentType": stored["ContentType"]}


def test_s3_store_keeps_original_encrypted_and_content_addressed():
    s3 = FakeS3()
    document = StatementDocument.sniff(PNG_BYTES)
    stored = S3DocumentStore("bucket-x", "ap-southeast-1", client=s3).put(document, "case-1")
    digest = document.ref.removeprefix("sha256:")
    assert stored == {
        "store": "s3",
        "bucket": "bucket-x",
        "key": f"statements/case-1/{digest}.png",
        "version_id": "v1",
    }
    call = s3.calls[0]
    assert call["ServerSideEncryption"] == "AES256" and call["ContentType"] == PNG
    assert call["Body"] == PNG_BYTES


def test_s3_store_fails_closed_and_needs_a_bucket():
    from botocore.exceptions import EndpointConnectionError

    broken = S3DocumentStore("b", "r", client=FakeS3(EndpointConnectionError(endpoint_url="x")))
    with pytest.raises(WorkflowError) as exc:
        broken.put(StatementDocument.sniff(PNG_BYTES), "c")
    assert (exc.value.code, exc.value.status) == ("DOCUMENT_STORE_UNAVAILABLE", 503)
    with pytest.raises(WorkflowError, match="DOCUMENT_STORE_NOT_CONFIGURED"):
        S3DocumentStore("", "r", client=FakeS3())
    assert factory.build_document_store("none").put(StatementDocument.sniff(PNG_BYTES), "c") == {
        "store": "none"
    }
    with pytest.raises(WorkflowError, match="DOCUMENT_STORE_UNKNOWN"):
        factory.build_document_store("ftp")


def test_uploaded_original_is_recorded_on_the_statement(client, upload_service):  # noqa: F811
    s3 = FakeS3()
    upload_service.documents = S3DocumentStore("bucket-x", "r", client=s3)
    case = upload_service.repository.get(create(client, "ducthinh.tech")["case_id"])
    _, _, (first_day, _) = period_rows(case)
    content = csv_bytes([[first_day, "SHOPEE VN THANH TOAN", "", "70.000.000"]])
    assert upload(client, case, content).status_code == 200
    statement = upload_service.repository.get(case.case_id).runs[-1]["statement"]
    assert statement["original"]["key"].startswith(f"statements/{case.case_id}/")
    assert isinstance(upload_service.documents, S3DocumentStore)
    assert NoDocumentStore().name == "none"
    content_back, media = upload_service.documents.get(statement["original"])
    assert content_back == content and media == "text/csv"


def test_local_store_round_trip_rejects_a_path_outside_the_root(tmp_path):
    store = LocalDocumentStore(tmp_path)
    document = StatementDocument.sniff(PNG_BYTES)
    record = store.put(document, "case-1")
    assert store.get(record) == (PNG_BYTES, document.media_type)
    with pytest.raises(WorkflowError) as escaped:
        store.get({**record, "path": "../outside"})
    assert escaped.value.code == "STATEMENT_FILE_NOT_STORED"
    with pytest.raises(WorkflowError) as missing:
        NoDocumentStore().get(record)
    assert missing.value.code == "STATEMENT_FILE_NOT_STORED"
    assert factory.build_document_store("local").name == "local"


def test_seller_view_locks_statement_period_from_payout_weeks(client, upload_service):  # noqa: F811
    waiting = upload_service.repository.get(create(client, "ducthinh.tech")["case_id"])
    weeks = waiting.runs[-1]["input_snapshot"]["features"]["weekly_payouts"]
    starts = [date.fromisoformat(row["week_start"]) for row in weeks]
    view = client.get(
        f"/v1/cases/{waiting.case_id}/seller-view", headers=seller_headers("ducthinh.tech")
    )
    assert view.json()["statement_period"] == {
        "start": min(starts).isoformat(),
        "end": (max(starts) + timedelta(days=7)).isoformat(),
    }
    ready = create(client)
    assert (
        client.get(f"/v1/cases/{ready['case_id']}/seller-view", headers=seller_headers()).json()[
            "statement_period"
        ]
        is None
    )


def test_statement_file_and_reading_can_be_downloaded(client, upload_service, tmp_path):  # noqa: F811
    upload_service.documents = LocalDocumentStore(tmp_path)
    case = upload_service.repository.get(create(client, "ducthinh.tech")["case_id"])
    missing = client.get(f"/v1/cases/{case.case_id}/statements/reading", headers=BANK)
    assert missing.status_code == 404
    assert missing.json()["detail"]["code"] == "STATEMENT_NOT_FOUND"
    _, _, (first_day, _) = period_rows(case)
    content = csv_bytes([[first_day, "SHOPEE VN THANH TOAN", "", "70.000.000"]])
    assert upload(client, case, content).status_code == 200
    original = client.get(f"/v1/cases/{case.case_id}/statements/original", headers=BANK)
    assert original.status_code == 200
    assert original.content == content
    assert "attachment" in original.headers["content-disposition"]
    assert original.headers["content-disposition"].endswith('-sao-ke.csv"')
    owner = seller_headers("ducthinh.tech")
    reading = client.get(f"/v1/cases/{case.case_id}/statements/reading", headers=owner)
    assert reading.status_code == 200
    body = reading.json()
    assert body["statement_lines"] and "original" not in body
    assert "ket-qua-doc.json" in reading.headers["content-disposition"]
    other = client.get(f"/v1/cases/{case.case_id}/statements/original", headers=seller_headers())
    assert other.status_code == 403
    view = client.get(f"/v1/cases/{case.case_id}/seller-view", headers=owner)
    assert view.json()["statement_file"]["stored"] is True


def test_unstored_original_still_returns_the_reading(client, upload_service):  # noqa: F811
    case = upload_service.repository.get(create(client, "ducthinh.tech")["case_id"])
    _, _, (first_day, _) = period_rows(case)
    content = csv_bytes([[first_day, "SHOPEE VN THANH TOAN", "", "70.000.000"]])
    assert upload(client, case, content).status_code == 200
    missing = client.get(f"/v1/cases/{case.case_id}/statements/original", headers=BANK)
    assert missing.status_code == 404
    assert missing.json()["detail"]["code"] == "STATEMENT_FILE_NOT_STORED"
    reading = client.get(f"/v1/cases/{case.case_id}/statements/reading", headers=BANK)
    assert reading.status_code == 200 and reading.json()["marketplace_total"] == 70_000_000
    view = client.get(
        f"/v1/cases/{case.case_id}/seller-view", headers=seller_headers("ducthinh.tech")
    )
    assert view.json()["statement_file"]["stored"] is False
