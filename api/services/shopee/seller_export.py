from __future__ import annotations

import csv
import re
from datetime import datetime, timezone
from pathlib import Path

from .schemas import EscrowDetail, Order, ReturnCase


class ColumnError(ValueError):
    pass


class SellerExportParser:
    ALIASES: dict[str, list[str]] = {
        "order_sn": ["mã đơn hàng", "order id", "order sn", "order_sn"],
        "create_time": ["ngày đặt hàng", "thời gian đặt hàng", "order creation date", "order time"],
        "order_status": ["trạng thái đơn hàng", "order status", "trạng thái"],
        "total_amount": [
            "tổng giá trị đơn hàng (vnd)",
            "tổng giá trị đơn hàng",
            "order total",
            "tổng tiền người mua trả",
            "tổng số tiền người mua thanh toán",
        ],
        "row_kind": ["đơn hàng/sản phẩm", "order/product"],
        "escrow_amount": [
            "tổng tiền đã thanh toán",
            "số tiền thanh toán",
            "tổng số tiền thanh toán",
            "total released amount",
            "thanh toán cuối cùng",
        ],
        "escrow_release_time": [
            "ngày hoàn thành thanh toán",
            "thời gian thanh toán",
            "ngày thanh toán",
            "payout completed date",
            "release time",
        ],
        "commission_fee": ["phí cố định", "commission fee"],
        "service_fee": ["phí dịch vụ", "service fee"],
        "seller_transaction_fee": ["phí xử lý giao dịch", "phí thanh toán", "transaction fee"],
        "return_sn": ["mã yêu cầu trả hàng", "mã yêu cầu", "return id", "return sn"],
        "return_time": ["ngày tạo", "thời gian tạo", "created time", "request time"],
        "return_status": ["trạng thái", "status"],
        "refund_amount": ["số tiền hoàn", "số tiền hoàn trả", "refund amount"],
    }
    STATUS_MAP: dict[str, str] = {
        "hoàn thành": "COMPLETED",
        "đã giao": "COMPLETED",
        "đang giao": "SHIPPED",
        "đã hủy": "CANCELLED",
        "chờ lấy hàng": "READY_TO_SHIP",
        "đang xử lý": "PROCESSED",
        "chờ xác nhận": "UNPAID",
    }
    INCOME_SHEETS = ("income", "doanh thu", "chi tiết doanh thu")
    ORDER_ROW_KINDS = ("đơn hàng", "order")
    TIME_FORMATS = ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d")

    def orders(self, path: str | Path) -> list[Order]:
        headers, rows = self._read(path)
        col = self._columns(
            headers, path, "order_sn", "create_time", "order_status", "total_amount"
        )
        return [
            Order(
                order_sn=str(r[col["order_sn"]]),
                create_time=self._ts(r[col["create_time"]]),
                order_status=self._status(r[col["order_status"]]),
                total_amount=self._money(r[col["total_amount"]]),
            )
            for r in rows
            if r[col["order_sn"]]
        ]

    def income(self, path: str | Path) -> list[EscrowDetail]:
        headers, rows = self._read(path, self.INCOME_SHEETS)
        col = self._columns(headers, path, "order_sn", "escrow_amount", "escrow_release_time")
        kind = self._optional(headers, "row_kind")
        if kind is not None:
            rows = [r for r in rows if self._norm(r[kind]) in self.ORDER_ROW_KINDS]
        opt = {
            k: self._optional(headers, k)
            for k in ("total_amount", "commission_fee", "service_fee", "seller_transaction_fee")
        }
        money = lambda r, k: self._money(r[opt[k]]) if opt[k] is not None else 0  # noqa: E731
        return [
            EscrowDetail(
                order_sn=str(r[col["order_sn"]]),
                escrow_amount=self._money(r[col["escrow_amount"]]),
                buyer_total_amount=money(r, "total_amount"),
                commission_fee=money(r, "commission_fee"),
                service_fee=money(r, "service_fee"),
                seller_transaction_fee=money(r, "seller_transaction_fee"),
                escrow_release_time=self._ts(r[col["escrow_release_time"]]),
            )
            for r in rows
            if r[col["order_sn"]]
        ]

    def returns(self, path: str | Path) -> list[ReturnCase]:
        headers, rows = self._read(path)
        col = self._columns(headers, path, "return_sn", "order_sn", "return_time", "return_status")
        refund = self._optional(headers, "refund_amount")
        return [
            ReturnCase(
                return_sn=str(r[col["return_sn"]]),
                order_sn=str(r[col["order_sn"]]),
                create_time=self._ts(r[col["return_time"]]),
                status=self._status(r[col["return_status"]]),
                refund_amount=self._money(r[refund]) if refund is not None else 0,
            )
            for r in rows
            if r[col["return_sn"]]
        ]

    def _read(
        self, path: str | Path, sheet_hint: tuple[str, ...] = ()
    ) -> tuple[list[str], list[list]]:
        p = Path(path)
        if p.suffix.lower() in {".xlsx", ".xlsm"}:
            rows = self._read_xlsx(p, sheet_hint)
        else:
            with open(p, encoding="utf-8-sig", newline="") as f:
                rows = list(csv.reader(f))
        rows = [r for r in rows if any(c not in (None, "") for c in r)]
        header_i = next(i for i, r in enumerate(rows) if sum(1 for c in r if c) >= 3)
        return [self._norm(c) for c in rows[header_i]], rows[header_i + 1 :]

    @staticmethod
    def _read_xlsx(p: Path, sheet_hint: tuple[str, ...]) -> list[list]:
        from openpyxl import load_workbook

        wb = load_workbook(p, read_only=True, data_only=True)
        ws = next(
            (wb[n] for n in wb.sheetnames if any(h in n.lower() for h in sheet_hint)), wb.active
        )
        return [list(r) for r in ws.iter_rows(values_only=True)]

    def _columns(self, headers: list[str], path, *keys: str) -> dict[str, int]:
        return {k: self._column(headers, k, str(path)) for k in keys}

    def _column(self, headers: list[str], key: str, path: str) -> int:
        aliases = self.ALIASES[key]
        for alias in aliases:
            if alias in headers:
                return headers.index(alias)
        for i, h in enumerate(headers):
            if any(alias in h for alias in aliases):
                return i
        raise ColumnError(f"{path}: không tìm thấy cột '{key}'. Cột có trong file: {headers}")

    def _optional(self, headers: list[str], key: str) -> int | None:
        try:
            return self._column(headers, key, "")
        except ColumnError:
            return None

    @staticmethod
    def _norm(s) -> str:
        return re.sub(r"\s+", " ", str(s or "").strip().lower())

    @staticmethod
    def _money(v) -> int:
        if v in (None, ""):
            return 0
        if isinstance(v, (int, float)):
            return int(round(v))
        digits = re.sub(r"[^\d\-]", "", str(v).replace(".", "").replace(",", ""))
        return int(digits) if digits else 0

    def _ts(self, v) -> int:
        if isinstance(v, datetime):
            return int(v.replace(tzinfo=v.tzinfo or timezone.utc).timestamp())
        text = str(v).strip()
        for fmt in self.TIME_FORMATS:
            try:
                return int(datetime.strptime(text, fmt).replace(tzinfo=timezone.utc).timestamp())
            except ValueError:
                continue
        raise ValueError(f"không đọc được thời gian: {v!r}")

    def _status(self, v) -> str:
        s = self._norm(v)
        for key, code in self.STATUS_MAP.items():
            if key in s:
                return code
        return s.upper().replace(" ", "_")
