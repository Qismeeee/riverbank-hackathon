import argparse
import io
import json
import random
import sys
import tempfile
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "api"))

import matplotlib  # noqa: E402
import pypdfium2  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font  # noqa: E402
from PIL import Image, ImageChops, ImageDraw, ImageEnhance, ImageFilter, ImageOps  # noqa: E402
from reportlab.lib import colors  # noqa: E402
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.lib.styles import ParagraphStyle  # noqa: E402
from reportlab.lib.units import mm  # noqa: E402
from reportlab.pdfbase import pdfmetrics  # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont  # noqa: E402
from reportlab.platypus import (  # noqa: E402
    Flowable,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from services.statements.classifiers import PayoutClassifier, PayoutLabel  # noqa: E402
from services.statements.factory import build_parsers  # noqa: E402
from services.statements.reader import StatementReader  # noqa: E402
from services.underwriting.cic import CicProfileBook, ProfileCicSource  # noqa: E402
from services.underwriting.gates import GateDecision, GateInput, OpenCicGate  # noqa: E402
from services.underwriting.models import Case, Plan, StatementUpload  # noqa: E402
from services.underwriting.repository import MemoryCaseRepository  # noqa: E402
from services.underwriting.risk import BaselineRiskModel, LightGbmRiskModel, RiskModel  # noqa: E402
from services.underwriting.service import UnderwritingService  # noqa: E402
from services.underwriting.sources import (  # noqa: E402
    InternalLosAdapter,
    SellerBook,
    SellerProfile,
    seller_book_source,
)

VN = timezone(timedelta(hours=7))
INK_HEX = "#1c2e86"
INK = colors.HexColor(INK_HEX)
INK_RGB = (28, 46, 134, 255)
SHADE = colors.HexColor("#e6e6e6")
PAPER = (246, 243, 235)
TIGHT = (
    ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ("TOPPADDING", (0, 0), (-1, -1), 2),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
)
BODY_FONTS = Path(matplotlib.get_data_path()) / "fonts/ttf"
PAYOUT_TAG = "SHOPEE VN IBFT thanh toan ky"
FEE_SHARES = {"fixed": 0.06, "payment": 0.03}
TAX_SHARES = {"vat": 0.01, "pit": 0.005}
SERVICE_SHARE = 0.025


def vnd(value: int | None) -> str:
    return f"{value:,}".replace(",", ".") if value else ""


def dmy(day: date) -> str:
    return day.strftime("%d/%m/%Y")


@dataclass(frozen=True)
class Window:
    now: datetime

    @property
    def monday(self) -> date:
        today = self.now.date()
        return today - timedelta(days=today.weekday())

    def week_start(self, index: int) -> date:
        return self.monday - timedelta(weeks=12 - index)

    @property
    def period(self) -> tuple[date, date]:
        return self.week_start(0), self.monday

    @property
    def statement_range(self) -> tuple[date, date]:
        return self.week_start(0) - timedelta(days=2), self.now.date() - timedelta(days=1)


@dataclass
class Entry:
    day: date
    desc: str
    debit: int = 0
    credit: int = 0
    payout: bool = False
    balance: int = 0


@dataclass
class Dossier:
    spec: dict
    seller: SellerProfile
    credit: dict
    window: Window
    lender: str = ""
    account_bank: str = ""
    disclosure: str = ""
    hands: tuple[str, ...] = ()
    entries: list[Entry] = field(default_factory=list)
    expected: dict = field(default_factory=dict)
    opening: int = 0

    def rng(self, salt: str) -> random.Random:
        return random.Random(f"{self.seller.seller_id}:{salt}")

    @property
    def city(self) -> str:
        return self.spec["business"]["address"].split(", ")[-1].removeprefix("TP. ")

    @property
    def birth_date(self) -> date:
        r = self.rng("birth")
        return date(self.spec["owner"]["birth_year"], r.randint(1, 12), r.randint(1, 28))

    @property
    def id_issued(self) -> date:
        r = self.rng("id")
        return date(r.randint(2021, 2024), r.randint(1, 12), r.randint(1, 28))

    @property
    def registered(self) -> date:
        month, year = self.spec["business"]["registered"].split("/")
        return date(int(year), int(month), self.rng("registered").randint(1, 28))

    @property
    def id_number(self) -> str:
        return f"********{self.spec['owner']['id_tail']}"

    @property
    def phone(self) -> str:
        return f"09****{self.spec['owner']['phone_tail']}"

    @property
    def account_number(self) -> str:
        return f"0123*****{self.spec['owner']['id_tail']}"

    def supplier_day(self, week: int) -> date:
        return self.window.monday + timedelta(weeks=week, days=2)

    @property
    def code(self) -> str:
        return self.spec["code"]

    @property
    def folder(self) -> str:
        return f"{self.code}-{self.seller.seller_id}"


class LedgerBuilder:
    def __init__(self, dossier: Dossier):
        self.d = dossier
        self.rng = random.Random(dossier.seller.seller_id)

    def build(self) -> list[Entry]:
        spec, seller, window = self.d.spec, self.d.seller, self.d.window
        mode = spec["statement"]
        share = mode.get("payout_share", 1.0)
        shop = spec["business"]["shop_name"]
        rows: list[Entry] = []
        for i, amount in enumerate(seller.weekly_payouts):
            start = window.week_start(i)
            paid = int(round(amount * share / 1000)) * 1000
            rows.append(
                Entry(
                    start + timedelta(days=1),
                    f"{PAYOUT_TAG} {start:%d/%m}-{start + timedelta(days=6):%d/%m} {shop}",
                    credit=paid,
                    payout=True,
                )
            )
            if share < 1:
                rows.append(self._filler(start, amount - paid, mode["fake_note"]))
            rows.extend(self._weekly_costs(start, amount))
        first, last = window.statement_range
        month = date(first.year, first.month, 1)
        while month <= last:
            rows.extend(self._monthly(month, first, last))
            month = date(month.year + month.month // 12, month.month % 12 + 1, 1)
        rows = sorted((r for r in rows if first <= r.day <= last), key=lambda r: (r.day, -r.credit))
        net = sum(r.credit - r.debit for r in rows)
        opening = spec["plan"]["cash_on_hand"] - net
        low = self._lowest(rows, opening)
        if low < 5_000_000:
            deposit = -(-(5_000_000 - low) // 1_000_000) * 1_000_000
            opening += deposit
            rows.append(
                Entry(last, f"Dat coc lo hang moi {spec['supplier']['name']}", debit=deposit)
            )
        balance = opening
        for r in rows:
            balance += r.credit - r.debit
            r.balance = balance
        self.d.opening = opening
        return rows

    def _filler(self, start: date, amount: int, note: str) -> Entry:
        names = ["NGUYEN VAN B", "TRAN THI C", "LE VAN D"]
        who = names[self.rng.randrange(len(names))]
        return Entry(
            start + timedelta(days=3), f"{who} {note}", credit=int(round(amount / 1000)) * 1000
        )

    def _weekly_costs(self, start: date, payout: int) -> list[Entry]:
        supplier = self.d.spec["supplier"]["name"]
        r = self.rng
        rows = [
            Entry(
                start + timedelta(days=2),
                f"CK thanh toan tien hang {supplier}",
                debit=int(payout * r.uniform(0.55, 0.7) / 1000) * 1000,
            ),
            Entry(
                start + timedelta(days=4),
                "Thanh toan phi van chuyen don hang",
                debit=int(payout * r.uniform(0.02, 0.04) / 1000) * 1000,
            ),
        ]
        if r.random() < 0.5:
            rows.append(
                Entry(
                    start + timedelta(days=5),
                    "Nap tien quang cao Shopee Ads",
                    debit=r.randrange(2, 6) * 1_000_000,
                )
            )
        if r.random() < 0.3:
            rows.append(
                Entry(
                    start + timedelta(days=6),
                    "Rut tien mat ATM",
                    debit=r.randrange(2, 8) * 1_000_000,
                )
            )
        return rows

    def _monthly(self, month: date, first: date, last: date) -> list[Entry]:
        r = self.rng
        rows = [
            Entry(
                month + timedelta(days=4),
                "Tra luong nhan vien dong goi",
                debit=r.randrange(14, 19) * 1_000_000,
            ),
            Entry(
                month + timedelta(days=9),
                "Thanh toan tien thue kho",
                debit=r.randrange(8, 13) * 1_000_000,
            ),
            Entry(
                month + timedelta(days=14),
                "Thanh toan tien dien nuoc internet",
                debit=r.randrange(1500, 3200) * 1000,
            ),
            Entry(month + timedelta(days=27), "Phi quan ly tai khoan va SMS", debit=22_000),
            Entry(
                month + timedelta(days=27),
                "Lai tien gui khong ky han",
                credit=r.randrange(8000, 40000),
            ),
        ]
        for f in self.d.credit.get("facilities", []):
            if "Thẻ" in f["product"]:
                rows.append(
                    Entry(
                        month + timedelta(days=19),
                        f"Thanh toan the tin dung {f['institution']}",
                        debit=max(1_000_000, f["outstanding"] // 10 // 1000 * 1000),
                    )
                )
            else:
                rows.append(
                    Entry(
                        month + timedelta(days=24),
                        f"Tra no vay {f['institution']}",
                        debit=max(2_000_000, f["outstanding"] // 12 // 1000 * 1000),
                    )
                )
        return [x for x in rows if first <= x.day <= last]

    @staticmethod
    def _lowest(rows: list[Entry], opening: int) -> int:
        low, balance = opening, opening
        for r in rows:
            balance += r.credit - r.debit
            low = min(low, balance)
        return low


class TruthClassifier(PayoutClassifier):
    name = "truth"
    model = "ground-truth"

    def classify(self, descriptions: list[str]) -> list[PayoutLabel]:
        return [PayoutLabel(PAYOUT_TAG.lower() in d.lower(), None) for d in descriptions]


class StatementFirstGate(OpenCicGate):
    model = "statement_match"
    waits_for_statement = True

    def evaluate(self, inp: GateInput) -> GateDecision:
        eligible = inp.statement is None or bool(inp.statement.get("matched"))
        return GateDecision(self.name, eligible, self.threshold, self.model)


class Expectation:
    def __init__(self, window: Window, risk: RiskModel):
        self.window, self.risk = window, risk

    def service(self) -> UnderwritingService:
        return UnderwritingService(
            MemoryCaseRepository(),
            seller_book_source(self.window.now),
            self.risk,
            ProfileCicSource(),
            InternalLosAdapter(),
            cic_gate=StatementFirstGate(),
            statements=StatementReader(build_parsers("off"), TruthClassifier()),
            clock=lambda: self.window.now,
        )

    def run(self, d: Dossier, statement_pdf: bytes) -> dict:
        svc = self.service()
        plan = Plan(seller_id=d.seller.seller_id, shop_id=d.seller.seller_id, **d.spec["plan"])
        case = svc.create(plan, f"dossier-{d.code}")
        start, end = self.window.period
        out = {"period": {"start": start.isoformat(), "end": end.isoformat()}}
        out["created"] = self._summary(case)
        if case.status == "AWAITING_STATEMENT":
            upload = StatementUpload(
                expected_version=case.version,
                period_start=start.isoformat(),
                period_end=end.isoformat(),
            )
            after = svc.upload_statement(case.case_id, upload, statement_pdf, d.seller.seller_id)
            out["after_statement"] = self._summary(after)
        return out

    @staticmethod
    def _summary(case: Case) -> dict:
        run = case.runs[-1]
        cf, risk, cic = run["cashflow"], run["risk"] or {}, run["cic"] or {}
        statement = run.get("statement") or {}
        failed = [r["rule_id"] for r in (run["policy"] or {}).get("results", []) if not r["passed"]]
        return {
            "status": case.status,
            "reason": case.reason,
            "gap": cf["gap"],
            "amount_proposed": cf.get("amount_proposed"),
            "tenor_weeks": cf.get("tenor_weeks"),
            "weekly_deduction": cf.get("weekly_deduction"),
            "payout_base": cf.get("payout_base"),
            "payout_warning_threshold": cf.get("payout_warning_threshold"),
            "pd_bps": risk.get("pd_bps"),
            "cic_group": cic.get("bureau_group"),
            "cic_worst_12m": cic.get("worst_group_12m"),
            "failed_rules": failed,
            "statement": {
                k: statement.get(k)
                for k in ("platform_total", "marketplace_total", "difference", "matched")
            }
            if statement
            else None,
            "cashflow_weeks": cf.get("weeks"),
        }


DIGITS = ["không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]
GROUPS = ["", " nghìn", " triệu", " tỷ"]


def _triple(n: int, full: bool) -> str:
    h, t, u = n // 100, n // 10 % 10, n % 10
    words = [DIGITS[h], "trăm"] if full or h else []
    if t == 0 and u and words:
        words.append("linh")
    elif t == 1:
        words.append("mười")
    elif t > 1:
        words += [DIGITS[t], "mươi"]
    if u:
        if u == 1 and t > 1:
            words.append("mốt")
        elif u == 5 and t > 0:
            words.append("lăm")
        elif u == 4 and t > 1:
            words.append("tư")
        else:
            words.append(DIGITS[u])
    return " ".join(words)


def in_words(amount: int) -> str:
    groups = []
    while amount:
        groups.append(amount % 1000)
        amount //= 1000
    if not groups:
        return "Không đồng"
    if len(groups) > len(GROUPS):
        raise ValueError("amount too large")
    parts = [
        _triple(g, i < len(groups) - 1) + GROUPS[i]
        for i, g in reversed(list(enumerate(groups)))
        if g
    ]
    text = " ".join(parts)
    return f"{text[0].upper()}{text[1:]} đồng"


@dataclass(frozen=True)
class Fonts:
    name: str
    bold: str
    hands: tuple[str, ...]

    @classmethod
    def register(cls, regular: Path, bold: Path, hand_dir: Path) -> "Fonts":
        pdfmetrics.registerFont(TTFont("Body", str(regular)))
        pdfmetrics.registerFont(TTFont("BodyBold", str(bold)))
        hands = []
        for n, path in enumerate(sorted(hand_dir.glob("*.ttf"))):
            pdfmetrics.registerFont(TTFont(f"Hand{n}", str(path)))
            hands.append(f"Hand{n}")
        if len(hands) < 2:
            raise SystemExit(f"cần ít nhất 2 font viết tay trong {hand_dir}")
        return cls("Body", "BodyBold", tuple(hands))

    def style(self, name: str, size: float = 8, bold: bool = False, **kw) -> ParagraphStyle:
        kw.setdefault("leading", size * 1.3)
        return ParagraphStyle(name, fontName=self.bold if bold else self.name, fontSize=size, **kw)


class Ink(ABC):
    def __init__(self, fonts: Fonts, style: ParagraphStyle):
        self.f, self.style = fonts, style

    @abstractmethod
    def inline(self, text: object) -> str: ...

    @abstractmethod
    def box(self, label: str, checked: bool) -> str: ...

    def value(self, text: object, align: int = TA_LEFT) -> Paragraph:
        return Paragraph(self.inline(text), ParagraphStyle("v", parent=self.style, alignment=align))

    def boxes(self, options: list[tuple[str, bool]]) -> Paragraph:
        return Paragraph("&nbsp;&nbsp;&nbsp; ".join(self.box(*o) for o in options), self.style)


class Typed(Ink):
    def __init__(self, fonts: Fonts):
        super().__init__(fonts, fonts.style("typed"))

    def inline(self, text: object) -> str:
        return escape(str(text))

    def box(self, label: str, checked: bool) -> str:
        return f"{'☒' if checked else '☐'} {escape(label)}"


class Handwritten(Ink):
    def __init__(self, fonts: Fonts, hand: str, rng: random.Random):
        super().__init__(fonts, fonts.style("hand", leading=13))
        self.hand, self.rng = hand, rng
        self.folder = tempfile.TemporaryDirectory()
        self.empty = self.mark("empty", False)
        self.ticks = [self.mark(f"tick{n}", True) for n in range(3)]

    def inline(self, text: object) -> str:
        size = round(self.rng.uniform(10, 11.5), 1)
        return (
            f'<font name="{self.hand}" size="{size}" color="{INK_HEX}">{escape(str(text))}</font>'
        )

    def box(self, label: str, checked: bool) -> str:
        src = self.rng.choice(self.ticks) if checked else self.empty
        return f'<img src="{src}" width="9" height="9" valign="-1"/> {escape(label)}'

    def mark(self, name: str, checked: bool) -> str:
        img = Image.new("RGBA", (64, 64), (0, 0, 0, 0))
        draw = ImageDraw.Draw(img)
        draw.rectangle((10, 10, 54, 54), outline=(0, 0, 0, 255), width=3)
        if checked:
            j = [self.rng.randint(-6, 6) for _ in range(8)]
            draw.line((6 + j[0], 8 + j[1], 58 + j[2], 56 + j[3]), fill=INK_RGB, width=6)
            draw.line((56 + j[4], 6 + j[5], 8 + j[6], 58 + j[7]), fill=INK_RGB, width=6)
        path = Path(self.folder.name) / f"{name}.png"
        img.save(path)
        return str(path)


class Signature(Flowable):
    def __init__(self, name: str, font: str, seed: str):
        super().__init__()
        self.text, self.font, self.rng = name.split()[-1], font, random.Random(seed)
        self.width, self.height = 50 * mm, 15 * mm

    def wrap(self, *_):
        return self.width, self.height

    def draw(self) -> None:
        c, r = self.canv, self.rng
        size = r.uniform(18, 23)
        span = c.stringWidth(self.text, self.font, size)
        c.saveState()
        c.setFillColor(INK)
        c.setStrokeColor(INK)
        c.setLineWidth(0.9)
        c.translate((self.width - span) / 2, self.height * 0.3)
        c.rotate(r.uniform(4, 11))
        c.setFont(self.font, size)
        c.drawString(0, 0, self.text)
        path = c.beginPath()
        path.moveTo(-4, -3)
        path.curveTo(span * 0.3, -9, span * 0.7, 4, span + 8, -2)
        c.drawPath(path, stroke=1, fill=0)
        c.restoreState()


class PdfDocument(ABC):
    key: str
    filename: str
    title: str
    label: str
    form_no = ""

    def __init__(self, d: Dossier, fonts: Fonts, ink: Ink):
        self.d, self.f, self.ink = d, fonts, ink

    @abstractmethod
    def story(self) -> list: ...

    def render(self) -> bytes:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=15 * mm,
            rightMargin=15 * mm,
            topMargin=12 * mm,
            bottomMargin=14 * mm,
            title=self.title.replace("<br/>", " "),
        )
        doc.build(
            self.header() + self.story(), onFirstPage=self.decorate, onLaterPages=self.decorate
        )
        return buffer.getvalue()

    def header(self) -> list:
        title = Paragraph(self.title, self.f.style("h1", 14, True, leading=18, alignment=TA_CENTER))
        return [self.letterhead(), Spacer(1, 8), title, *self.subtitle(), Spacer(1, 6)]

    def subtitle(self) -> list:
        return []

    def letterhead(self) -> Table:
        left = Paragraph(
            f"<b>{escape(self.d.lender.upper())}</b><br/>Chi nhánh {escape(self.d.city)}",
            self.f.style("lh", 8.5),
        )
        right = Paragraph(
            f"Mẫu số {self.form_no}" if self.form_no else "",
            self.f.style("fn", 8, alignment=TA_RIGHT),
        )
        return self.bare([[left, right]], [110, 70])

    def decorate(self, canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont(self.f.name, 8)
        canvas.setFillColor(colors.grey)
        canvas.drawRightString(A4[0] - 15 * mm, 8 * mm, f"Trang {doc.page}")
        canvas.setFillColor(colors.Color(0.8, 0.1, 0.1, alpha=0.06))
        canvas.setFont(self.f.bold, 64)
        canvas.translate(A4[0] / 2, A4[1] / 2)
        canvas.rotate(35)
        canvas.drawCentredString(0, 0, "MẪU")
        canvas.restoreState()

    def cell(self, text: object, bold: bool = False, align: int = TA_LEFT) -> Paragraph:
        return Paragraph(escape(str(text)), self.f.style("cell", bold=bold, alignment=align))

    def subtitle_line(self, text: str) -> Paragraph:
        return Paragraph(text, self.f.style("sub", 9, alignment=TA_CENTER))

    def heading(self, text: str) -> Paragraph:
        return Paragraph(escape(text), self.f.style("h2", 9.5, True, spaceBefore=8, spaceAfter=3))

    @staticmethod
    def bare(rows: list[list], widths: list[float]) -> Table:
        t = Table(rows, colWidths=[w * mm for w in widths])
        t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return t

    def grid(self, rows: list[list], widths: list[float], right=(), extra=()) -> Table:
        data = [
            [
                c
                if isinstance(c, Flowable)
                else self.cell(c, bold=i == 0, align=TA_CENTER if i == 0 else TA_LEFT)
                for c in row
            ]
            for i, row in enumerate(rows)
        ]
        style = [
            *TIGHT,
            ("BACKGROUND", (0, 0), (-1, 0), SHADE),
            *extra,
        ]
        style += [("ALIGN", (col, 1), (col, -1), "RIGHT") for col in right]
        t = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
        t.setStyle(TableStyle(style))
        return t

    def filled(self, value: object, align: int = TA_LEFT) -> Flowable:
        return value if isinstance(value, Flowable) else self.ink.value(value, align)

    def form(self, sections: list[tuple[str | None, list[tuple]]]) -> Table:
        rows, style = [], []
        for heading, items in sections:
            if heading:
                r = len(rows)
                style += [("SPAN", (0, r), (-1, r)), ("BACKGROUND", (0, r), (-1, r), SHADE)]
                rows.append([self.cell(heading, bold=True), "", "", ""])
            for item in items:
                r = len(rows)
                if len(item) == 1:
                    rows.append([self.filled(item[0]), "", "", ""])
                    style.append(("SPAN", (0, r), (-1, r)))
                elif len(item) == 2:
                    rows.append([self.cell(item[0]), self.filled(item[1]), "", ""])
                    style.append(("SPAN", (1, r), (-1, r)))
                else:
                    rows.append(
                        [
                            self.cell(item[0]),
                            self.filled(item[1]),
                            self.cell(item[2]),
                            self.filled(item[3]),
                        ]
                    )
        t = Table(rows, colWidths=[w * mm for w in (38, 56, 30, 56)])
        t.setStyle(
            TableStyle(
                [
                    *TIGHT,
                    *style,
                ]
            )
        )
        return t

    def place_date(self, day: date) -> Paragraph:
        ink = self.ink.inline
        return Paragraph(
            f"{escape(self.d.city)}, ngày {ink(f'{day:%d}')} tháng {ink(f'{day:%m}')} "
            f"năm {ink(day.year)}",
            self.f.style("pd", 9, alignment=TA_RIGHT, leading=16),
        )

    def signatures(self, parties: list[tuple[str, str | None, str | None]]) -> Table:
        center = self.f.style("sg", 8.5, alignment=TA_CENTER)
        width = 180 / len(parties)
        titles, marks, names = [], [], []
        for title, name, font in parties:
            titles.append(Paragraph(f"<b>{escape(title)}</b><br/>(Ký, ghi rõ họ tên)", center))
            marks.append(
                Signature(name, font, f"{self.d.code}:{self.key}:{name}") if name and font else ""
            )
            names.append(self.ink.value(name, TA_CENTER) if name else "")
        t = Table([titles, marks, names], colWidths=[width * mm] * len(parties))
        t.setStyle(
            TableStyle(
                [("ALIGN", (0, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
            )
        )
        return t


CONSENTS = (
    "Đồng ý cho Ngân hàng tra cứu thông tin tín dụng",
    "Đồng ý cho Ngân hàng nhận dữ liệu doanh thu, thanh toán từ sàn",
    "Thông tin kê khai trên đây là đúng sự thật",
)
DELIVERY = "Giao tại địa điểm kinh doanh của Bên B trong 03 ngày kể từ đợt thanh toán đầu tiên."


class LoanApplication(PdfDocument):
    key = "01"
    filename = "01_giay-de-nghi-vay-kiem-phuong-an-kinh-doanh.pdf"
    title = "GIẤY ĐỀ NGHỊ VAY VỐN<br/>KIÊM PHƯƠNG ÁN KINH DOANH"
    label = "Giấy đề nghị vay vốn kiêm phương án kinh doanh"
    form_no = "01/HKD"

    def story(self) -> list:
        s, plan, w, ink = self.d.spec, self.d.spec["plan"], self.d.window, self.ink
        o, b, sup = s["owner"], s["business"], s["supplier"]
        created = self.d.expected["created"]
        amount = plan["requested_amount"]
        schedule = ", ".join(
            f"{vnd(x['amount'])} đ ngày {dmy(self.d.supplier_day(x['week']))}"
            for x in plan["supplier_schedule"]
        )
        story = [
            self.form(
                [
                    (
                        "I. THÔNG TIN KHÁCH HÀNG",
                        [
                            ("Họ và tên", o["name"], "Ngày sinh", dmy(self.d.birth_date)),
                            ("Số CCCD", self.d.id_number, "Điện thoại", self.d.phone),
                            ("Địa chỉ cư trú", b["address"]),
                            ("Tên hộ kinh doanh", b["name"]),
                            ("Ngành nghề", b["category"]),
                            (
                                "Gian hàng TMĐT",
                                f"Shopee: {b['shop_name']}",
                                "Thời gian bán",
                                f"{self.d.seller.profile['shop_age_months']} tháng",
                            ),
                        ],
                    ),
                    (
                        "II. ĐỀ NGHỊ VAY VỐN",
                        [
                            ("Số tiền vay", f"{vnd(amount)} đồng"),
                            ("Bằng chữ", in_words(amount)),
                            ("Mục đích", f"Bổ sung vốn lưu động nhập hàng: {sup['goods']}"),
                            ("Thời hạn vay", f"{plan['sell_weeks']} tuần"),
                            (
                                "Phương thức trả nợ",
                                ink.boxes(
                                    [
                                        ("Trích tiền sàn hằng tuần", True),
                                        ("Trả góp hằng tháng", False),
                                        ("Trả gốc cuối kỳ", False),
                                    ]
                                ),
                            ),
                            ("Tài sản bảo đảm", ink.boxes([("Không", True), ("Có", False)])),
                            ("Giải ngân", f"Chuyển khoản cho {sup['name']}"),
                        ],
                    ),
                    (
                        "III. PHƯƠNG ÁN KINH DOANH",
                        [
                            (
                                "Tổng nhu cầu vốn",
                                f"{vnd(plan['lot_value'])} đồng",
                                "Vốn tự có",
                                f"{vnd(plan['cash_on_hand'])} đồng",
                            ),
                            (
                                "Hàng hóa",
                                f"{sup['goods']}, {vnd(sup['units'])} sản phẩm x "
                                f"{vnd(sup['unit_price'])} đồng",
                            ),
                            ("Nhà cung cấp", sup["name"]),
                            ("Lịch trả nhà cung cấp", schedule),
                            (
                                "Tiền sàn về bình quân",
                                f"{vnd(round(created['payout_base'], -6))} đồng mỗi tuần",
                            ),
                        ],
                    ),
                ]
            ),
        ]
        weeks = created.get("cashflow_weeks") or []
        if weeks:
            rows = [["Tuần", "Từ ngày", "Trả nhà cung cấp", "Tiền sàn về", "Số dư cuối tuần"]]
            for week in weeks:
                rows.append(
                    [
                        self.cell(week["week"] + 1, align=TA_CENTER),
                        self.cell(dmy(w.monday + timedelta(weeks=week["week"]))),
                        ink.value(vnd(week["outflow"]), TA_RIGHT),
                        ink.value(vnd(week["payout"]), TA_RIGHT),
                        ink.value(vnd(week["bal_end"]), TA_RIGHT),
                    ]
                )
            story += [self.heading("IV. KẾ HOẠCH DÒNG TIỀN"), self.grid(rows, [16, 28, 45, 45, 46])]
        story += [
            Spacer(1, 6),
            KeepTogether(
                [
                    self.form(
                        [
                            (
                                "V. CAM KẾT CỦA KHÁCH HÀNG",
                                [(ink.boxes([(text, True)]),) for text in CONSENTS],
                            )
                        ]
                    ),
                    Spacer(1, 8),
                    self.place_date(w.now.date()),
                    self.signatures(
                        [
                            ("CÁN BỘ TIẾP NHẬN", None, None),
                            ("KHÁCH HÀNG", o["name"], self.d.hands[0]),
                        ]
                    ),
                ]
            ),
        ]
        return story


class CustomerProfile(PdfDocument):
    key = "02"
    filename = "02_phieu-thong-tin-khach-hang.pdf"
    title = "PHIẾU THÔNG TIN KHÁCH HÀNG<br/>HỘ KINH DOANH"
    label = "Phiếu thông tin khách hàng hộ kinh doanh"
    form_no = "02/KYC"

    def story(self) -> list:
        s, ink, d = self.d.spec, self.ink, self.d
        o, b = s["owner"], s["business"]
        has_loans = bool(d.credit.get("facilities"))
        return [
            self.form(
                [
                    (
                        "I. CHỦ HỘ KINH DOANH",
                        [
                            ("Họ và tên", o["name"], "Ngày sinh", dmy(d.birth_date)),
                            ("Quốc tịch", "Việt Nam", "Số CCCD", d.id_number),
                            ("Ngày cấp", dmy(d.id_issued), "Điện thoại", d.phone),
                            ("Địa chỉ cư trú", b["address"]),
                            ("Nghề nghiệp", "Kinh doanh"),
                        ],
                    ),
                    (
                        "II. HỘ KINH DOANH",
                        [
                            ("Tên hộ kinh doanh", b["name"]),
                            ("Mã số thuế", d.id_number, "Ngày đăng ký", dmy(d.registered)),
                            ("Địa điểm kinh doanh", b["address"]),
                            ("Ngành nghề", b["category"]),
                            ("Kênh bán", f"Shopee, gian hàng {b['shop_name']}"),
                        ],
                    ),
                    (
                        "III. TÀI KHOẢN NHẬN TIỀN",
                        [
                            ("Ngân hàng", d.account_bank),
                            ("Số tài khoản", d.account_number, "Chủ tài khoản", o["name"].upper()),
                        ],
                    ),
                    (
                        "IV. THÔNG TIN KHÁC",
                        [
                            (
                                "Đang vay tại tổ chức khác",
                                ink.boxes([("Có", has_loans), ("Không", not has_loans)]),
                            ),
                            ("Cư trú tại Hoa Kỳ", ink.boxes([("Có", False), ("Không", True)])),
                            (
                                "Sản phẩm đăng ký",
                                ink.boxes(
                                    [("Vay vốn kinh doanh", True), ("Tài khoản thanh toán", True)]
                                ),
                            ),
                        ],
                    ),
                ]
            ),
            Spacer(1, 10),
            self.place_date(d.window.now.date()),
            self.signatures(
                [("CÁN BỘ NGÂN HÀNG", None, None), ("KHÁCH HÀNG", o["name"], d.hands[0])]
            ),
        ]


class StatementPdf(PdfDocument):
    key = "03"
    filename = "03_sao-ke-tai-khoan.pdf"
    title = "SAO KÊ TÀI KHOẢN THANH TOÁN"
    label = "Sao kê tài khoản nhận tiền sàn"

    def story(self) -> list:
        first, last = self.d.window.statement_range
        o = self.d.spec["owner"]
        rows = [["STT", "Ngày giao dịch", "Nội dung", "Ghi nợ", "Ghi có", "Số dư"]]
        rows.append(["", "", "Số dư đầu kỳ", "", "", vnd(self.d.opening)])
        for n, e in enumerate(self.d.entries, start=1):
            rows.append([n, dmy(e.day), e.desc, vnd(e.debit), vnd(e.credit), vnd(e.balance)])
        rows.append(["", "", "Số dư cuối kỳ", "", "", vnd(self.d.entries[-1].balance)])
        return [
            self.form(
                [
                    (
                        None,
                        [
                            (
                                "Chủ tài khoản",
                                o["name"].upper(),
                                "Số tài khoản",
                                self.d.account_number,
                            ),
                            ("Từ ngày", dmy(first), "Đến ngày", dmy(last)),
                            ("Loại tiền", "VND", "Ngày in", dmy(self.d.window.now.date())),
                        ],
                    )
                ]
            ),
            Spacer(1, 6),
            self.grid(rows, [10, 22, 76, 24, 24, 26], right=(3, 4, 5)),
        ]


class PurchaseContract(PdfDocument):
    key = "05"
    filename = "05_hop-dong-mua-ban-hang-hoa.pdf"
    title = "HỢP ĐỒNG MUA BÁN HÀNG HÓA"
    label = "Hợp đồng mua bán hàng hóa"

    def letterhead(self) -> Paragraph:
        return Paragraph(
            "<b>CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</b><br/>Độc lập - Tự do - Hạnh phúc",
            self.f.style("motto", 10, alignment=TA_CENTER, leading=14),
        )

    def subtitle(self) -> list:
        number = self.d.rng("contract").randrange(10, 99)
        return [self.subtitle_line(f"Số: {number}/{self.signed.year}/HĐMB")]

    @property
    def signed(self) -> date:
        return self.d.window.monday - timedelta(days=3)

    def story(self) -> list:
        s, d = self.d.spec, self.d
        sup, plan, b, o = s["supplier"], s["plan"], s["business"], s["owner"]
        pay = [["Đợt", "Ngày thanh toán", "Số tiền (đồng)", "Hóa đơn"]]
        for n, x in enumerate(plan["supplier_schedule"], start=1):
            pay.append(
                [n, dmy(d.supplier_day(x["week"])), vnd(x["amount"]), f"Hóa đơn GTGT đợt {n}"]
            )
        goods = [
            ["STT", "Tên hàng", "Số lượng", "Đơn giá (đồng)", "Thành tiền (đồng)"],
            [1, sup["goods"], vnd(sup["units"]), vnd(sup["unit_price"]), vnd(plan["lot_value"])],
            ["", "Tổng cộng (đã gồm thuế GTGT)", "", "", vnd(plan["lot_value"])],
        ]
        body = self.f.style("cb", 9, leading=13)
        return [
            Paragraph(f"Hôm nay, ngày {dmy(self.signed)}, chúng tôi gồm:", body),
            Spacer(1, 4),
            self.form(
                [
                    (
                        "BÊN BÁN (BÊN A)",
                        [
                            ("Tên", sup["name"]),
                            ("Địa chỉ", sup["address"]),
                            ("Đại diện", sup["rep"], "Chức vụ", "Giám đốc"),
                        ],
                    ),
                    (
                        "BÊN MUA (BÊN B)",
                        [
                            ("Tên", b["name"]),
                            ("Địa chỉ", b["address"]),
                            ("Đại diện", o["name"], "Chức vụ", "Chủ hộ kinh doanh"),
                        ],
                    ),
                ]
            ),
            self.heading("Điều 1. Hàng hóa"),
            self.grid(
                goods,
                [12, 78, 24, 30, 36],
                right=(2, 3, 4),
                extra=[("SPAN", (1, 2), (3, 2))],
            ),
            self.heading("Điều 2. Thanh toán"),
            self.grid(pay, [14, 40, 46, 80], right=(2,)),
            self.heading("Điều 3. Giao nhận"),
            Paragraph(
                DELIVERY,
                body,
            ),
            self.heading("Điều 4. Hóa đơn"),
            Paragraph("Bên A xuất hóa đơn điện tử cho Bên B theo từng đợt thanh toán.", body),
            Spacer(1, 10),
            self.signatures(
                [
                    ("ĐẠI DIỆN BÊN A", sup["rep"], d.hands[1]),
                    ("ĐẠI DIỆN BÊN B", o["name"], d.hands[0]),
                ]
            ),
        ]


class CreditReport(PdfDocument):
    key = "07"
    filename = "07_ket-qua-tra-cuu-thong-tin-tin-dung.pdf"
    title = "KẾT QUẢ TRA CỨU THÔNG TIN TÍN DỤNG"
    label = "Kết quả tra cứu thông tin tín dụng"

    def subtitle(self) -> list:
        ref = self.d.rng("cic").randrange(10**7, 10**8)
        return [self.subtitle_line(f"Mã tra cứu: {ref}")]

    def story(self) -> list:
        o, c, d = self.d.spec["owner"], self.d.credit, self.d
        facilities = c.get("facilities", [])
        total = sum(f["outstanding"] for f in facilities)
        rows = [["Tổ chức tín dụng", "Sản phẩm", "Hạn mức", "Dư nợ", "Nhóm nợ", "Thời gian"]]
        for f in facilities:
            rows.append(
                [
                    f["institution"],
                    f["product"],
                    vnd(f["limit"]),
                    vnd(f["outstanding"]),
                    f["group"],
                    f"{f['opened_months_ago']} tháng",
                ]
            )
        extra = []
        if facilities:
            rows.append(["Tổng cộng", "", "", vnd(total), "", ""])
            extra = [("SPAN", (0, len(rows) - 1), (2, len(rows) - 1))]
        else:
            rows.append(["Không có quan hệ tín dụng", "", "", "", "", ""])
            extra = [("SPAN", (0, 1), (-1, 1))]
        first = d.window.now.date().replace(day=1)
        months = [
            (first - timedelta(days=30 * k)).replace(day=1).strftime("%m/%y")
            for k in range(12, 0, -1)
        ]
        history = [["Khoản vay", *months]]
        history += [[f["product"], *f["history"]] for f in facilities if f.get("history")]
        if len(history) == 1:
            history.append(["Không có dữ liệu", *[""] * 12])
        groups = [g for f in facilities for g in f.get("history", [])]
        return [
            self.form(
                [
                    (
                        None,
                        [
                            ("Khách hàng", o["name"], "Số CCCD", d.id_number),
                            ("Ngày tra cứu", dmy(d.window.now.date()), "Đơn vị tra cứu", d.lender),
                            ("Số lần bị tra cứu 90 ngày", str(c.get("recent_inquiries_90d", 0))),
                        ],
                    )
                ]
            ),
            self.heading("1. Dư nợ hiện tại"),
            self.grid(rows, [36, 36, 28, 28, 22, 30], right=(2, 3), extra=extra),
            self.heading("2. Nhóm nợ 12 tháng gần nhất"),
            self.grid(history, [18] + [13.5] * 12),
            self.heading("3. Tổng hợp"),
            self.form(
                [
                    (
                        None,
                        [
                            (
                                "Nhóm nợ hiện tại",
                                str(max((f["group"] for f in facilities), default="Không có")),
                                "Nhóm cao nhất 12 tháng",
                                str(max(groups, default="Không có")),
                            ),
                            ("Tổng dư nợ", f"{vnd(total) or 0} đồng"),
                        ],
                    )
                ]
            ),
        ]


class Paper(ABC):
    note = ""

    @abstractmethod
    def apply(self, pdf: bytes, rng: random.Random) -> bytes: ...


class Digital(Paper):
    def apply(self, pdf: bytes, rng: random.Random) -> bytes:
        return pdf


@dataclass(frozen=True)
class ScanProfile:
    scale: float
    tilt: float
    blur: float
    noise: float
    quality: int
    gray: bool
    shadow: float


SCAN_PROFILES = {
    "normal": ScanProfile(2.0, 0.9, 0.4, 0.05, 72, False, 0.0),
    "poor": ScanProfile(1.5, 2.0, 0.8, 0.1, 40, True, 0.45),
}


class Scanned(Paper):
    note = "bản scan"

    def __init__(self, profile: ScanProfile):
        self.p = profile

    def apply(self, pdf: bytes, rng: random.Random) -> bytes:
        pages = [
            self.page(p.render(scale=self.p.scale).to_pil(), rng)
            for p in pypdfium2.PdfDocument(pdf)
        ]
        buffer = io.BytesIO()
        pages[0].save(
            buffer, "PDF", save_all=True, append_images=pages[1:], resolution=72 * self.p.scale
        )
        return buffer.getvalue()

    def page(self, img: Image.Image, rng: random.Random) -> Image.Image:
        p = self.p
        img = ImageChops.multiply(img.convert("RGB"), Image.new("RGB", img.size, PAPER))
        if p.gray:
            img = ImageEnhance.Contrast(ImageOps.grayscale(img)).enhance(1.5).convert("RGB")
        img = img.rotate(rng.uniform(-p.tilt, p.tilt), resample=Image.BICUBIC, fillcolor=PAPER)
        if p.shadow:
            ramp = Image.linear_gradient("L").rotate(90).resize(img.size)
            mask = ramp.point(lambda v: int(max(0, v - 200) * 4.6 * p.shadow))
            img = Image.composite(Image.new("RGB", img.size, (40, 40, 40)), img, mask)
        draw = ImageDraw.Draw(img)
        for _ in range(int(60 * p.noise * 10)):
            x, y = rng.randrange(img.width), rng.randrange(img.height)
            r = rng.choice((1, 1, 2))
            draw.ellipse((x, y, x + r, y + r), fill=(70, 70, 70))
        img = img.filter(ImageFilter.GaussianBlur(p.blur))
        img = Image.blend(img, Image.effect_noise(img.size, 40).convert("RGB"), p.noise)
        buffer = io.BytesIO()
        img.save(buffer, "JPEG", quality=p.quality)
        return Image.open(io.BytesIO(buffer.getvalue()))


class Handfilled(Scanned):
    note = "viết tay, bản scan"


def paper_for(mode: str, quality: str) -> Paper:
    if mode == "digital":
        return Digital()
    profile = SCAN_PROFILES[quality]
    return {"scan": Scanned, "hand": Handfilled}[mode](profile)


class Workbook_(ABC):
    key: str
    filename: str
    label: str

    def __init__(self, d: Dossier):
        self.d = d

    @abstractmethod
    def fill(self, wb: Workbook) -> None: ...

    def render(self) -> bytes:
        wb = Workbook()
        self.fill(wb)
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def bold_row(ws, row: list) -> None:
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)


def platform_rows(d: Dossier) -> list[dict]:
    rows = []
    for i, payout in enumerate(d.seller.weekly_payouts):
        start = d.window.week_start(i)
        net_share = 1 - sum(FEE_SHARES.values()) - sum(TAX_SHARES.values()) - SERVICE_SHARE
        gross = round(payout / net_share / 1000) * 1000
        fees = {k: round(gross * v) for k, v in FEE_SHARES.items()}
        taxes = {k: round(gross * v) for k, v in TAX_SHARES.items()}
        service = gross - payout - sum(fees.values()) - sum(taxes.values())
        rows.append(
            {
                "start": start,
                "end": start + timedelta(days=6),
                "gross": gross,
                **fees,
                "service": service,
                **taxes,
                "payout": payout,
            }
        )
    return rows


class PlatformReport(Workbook_):
    key = "04"
    filename = "04_bao-cao-thanh-toan-san.xlsx"
    label = "Báo cáo thanh toán của sàn theo tuần"

    def fill(self, wb: Workbook) -> None:
        ws = wb.active
        ws.title = "Thanh toan"
        first, last = self.d.window.period
        self.bold_row(ws, ["BÁO CÁO THANH TOÁN"])
        ws.append(["Tên shop", self.d.spec["business"]["shop_name"]])
        ws.append(["Kỳ báo cáo", f"{dmy(first)} đến {dmy(last - timedelta(days=1))}"])
        ws.append([])
        self.bold_row(
            ws,
            [
                "Từ ngày",
                "Đến ngày",
                "Doanh thu đơn hàng",
                "Phí cố định",
                "Phí thanh toán",
                "Phí dịch vụ",
                "Thuế GTGT khấu trừ 1%",
                "Thuế TNCN khấu trừ 0,5%",
                "Số tiền thanh toán về tài khoản",
            ],
        )
        rows = platform_rows(self.d)
        keys = ("gross", "fixed", "payment", "service", "vat", "pit", "payout")
        for r in rows:
            ws.append([dmy(r["start"]), dmy(r["end"]), *(r[k] for k in keys)])
        self.bold_row(ws, ["Tổng cộng", "", *(sum(r[k] for r in rows) for k in keys)])
        for col in "ABCDEFGHI":
            ws.column_dimensions[col].width = 18


class RevenueBook(Workbook_):
    key = "06"
    filename = "06_so-doanh-thu-S1a-HKD.xlsx"
    label = "Sổ doanh thu bán hàng mẫu S1a-HKD"

    def fill(self, wb: Workbook) -> None:
        ws = wb.active
        ws.title = "S1a-HKD"
        b = self.d.spec["business"]
        ws.append([f"HỘ, CÁ NHÂN KINH DOANH: {b['name']}", "", "", "", "Mẫu số S1a-HKD"])
        ws.append([f"Địa chỉ: {b['address']}"])
        ws.append([])
        self.bold_row(ws, ["", "", "SỔ CHI TIẾT DOANH THU BÁN HÀNG HÓA, DỊCH VỤ"])
        ws.append(["", "", f"Địa điểm kinh doanh: {b['address']}"])
        ws.append(["", "", f"Kỳ: {self.d.window.now.year}"])
        ws.append([])
        self.bold_row(
            ws, ["Ngày ghi sổ", "Số hiệu chứng từ", "Ngày chứng từ", "Diễn giải", "Số tiền"]
        )
        month, subtotal, total = None, 0, 0
        for n, r in enumerate(platform_rows(self.d), start=1):
            if month and r["end"].month != month:
                self.bold_row(ws, ["", "", "", f"Cộng tháng {month:02d}", subtotal])
                subtotal = 0
            month = r["end"].month
            ws.append(
                [
                    dmy(r["end"]),
                    f"BCTT{n:03d}",
                    dmy(r["end"]),
                    f"Doanh thu bán hàng qua sàn từ {r['start']:%d/%m} đến {r['end']:%d/%m}",
                    r["gross"],
                ]
            )
            subtotal += r["gross"]
            total += r["gross"]
        self.bold_row(ws, ["", "", "", f"Cộng tháng {month:02d}", subtotal])
        self.bold_row(ws, ["", "", "", "Tổng cộng", total])
        ws.append([])
        ws.append(["", "", "", f"Ngày {dmy(self.d.window.now.date())}"])
        self.bold_row(ws, ["", "", "", "NGƯỜI ĐẠI DIỆN HỘ KINH DOANH"])
        ws.append(["", "", "", self.d.spec["owner"]["name"]])
        for col, width in zip("ABCDE", (14, 16, 14, 52, 18), strict=True):
            ws.column_dimensions[col].width = width


class StatementXlsx(Workbook_):
    key = "03b"
    filename = "03b_sao-ke-tai-khoan.xlsx"
    label = "Sao kê tài khoản bản Excel"

    def fill(self, wb: Workbook) -> None:
        ws = wb.active
        ws.title = "Sao ke"
        self.bold_row(ws, ["Ngày giao dịch", "Nội dung", "Ghi nợ", "Ghi có", "Số dư"])
        for e in self.d.entries:
            ws.append([dmy(e.day), e.desc, e.debit or None, e.credit or None, e.balance])


def photo(pdf: bytes) -> bytes:
    page = pypdfium2.PdfDocument(pdf)[0].render(scale=2.2).to_pil().convert("RGB")
    page = page.rotate(1.4, expand=True, fillcolor=(96, 90, 80), resample=Image.BICUBIC)
    page = page.filter(ImageFilter.GaussianBlur(0.9))
    page = Image.blend(page, Image.effect_noise(page.size, 20).convert("RGB"), 0.05)
    buffer = io.BytesIO()
    page.save(buffer, "JPEG", quality=78)
    return buffer.getvalue()


PDF_DOCUMENTS = (LoanApplication, CustomerProfile, PurchaseContract, CreditReport)
WORKBOOKS = (PlatformReport, RevenueBook, StatementXlsx)
PHOTO = ("03", "03_sao-ke-anh-chup.jpg", "Ảnh chụp trang 1 sao kê")
SCAN = ("03c", "03c_sao-ke-ban-scan.pdf", "Sao kê bản scan, dùng thử OCR")


class Readme:
    filename = "README.md"

    def __init__(self, d: Dossier, docs: list[tuple[str, str, str]]):
        self.d, self.docs = d, docs

    def render(self) -> bytes:
        s, w, e = self.d.spec, self.d.window, self.d.expected
        start, end = w.period
        lines = [
            f"# {s['code']} · {s['title']}",
            "",
            f"> {self.d.disclosure}",
            "",
            s["story"],
            "",
            "## Đăng nhập",
            "",
            f"Seller: `{self.d.seller.seller_id}`. Ngày hồ sơ: {dmy(w.now.date())}. "
            "Ngày khác thì chạy lại `make dossiers DATE=yyyy-mm-dd` "
            "để ngày trên sao kê khớp với payout.",
            "",
            "## Câu chat hoặc form",
            "",
            f"> {s['chat']}",
            "",
            "| Trường | Giá trị |",
            "|---|---|",
            *[f"| `{k}` | {json.dumps(v, ensure_ascii=False)} |" for k, v in s["plan"].items()],
            "",
            "## Kết quả Riverbank phải ra",
            "",
            *self.result("Sau khi tạo hồ sơ", e["created"]),
        ]
        if "after_statement" in e:
            lines += [
                "",
                f"Nộp `{StatementPdf.filename}` với `period_start={start}` và "
                f"`period_end={end}`. Bản `{SCAN[1]}` đủ các trang, cần bật OCR, "
                f"kết quả phải giống bản PDF. Ảnh `{PHOTO[1]}` chỉ có trang 1.",
                "",
                *self.result("Sau khi nộp sao kê", e["after_statement"]),
            ]
            if s["statement"]["mode"] == "inflated":
                lines += [
                    "",
                    "Kết quả trên giả định Jev loại đúng các dòng chuyển khoản cá nhân. "
                    "Nếu chạy `RIVERBANK_STATEMENT_CLASSIFIER=pattern` thì các dòng đó lọt qua vì "
                    "có chữ shopee, hồ sơ sẽ khớp. Đây chính là lý do dùng Jev.",
                ]
        lines += [
            "",
            "## Kịch bản trình bày",
            "",
            *[f"{n}. {x}" for n, x in enumerate(s["steps"], 1)],
            "",
            "## Tài liệu",
            "",
            "| File | Loại giấy | Dạng |",
            "|---|---|---|",
            *[f"| `{name}` | {label} | {form} |" for name, label, form in self.docs],
            "",
        ]
        return "\n".join(lines).encode("utf-8")

    @staticmethod
    def result(label: str, r: dict) -> list[str]:
        rows = [
            ("Trạng thái", f"`{r['status']}`" + (f" ({r['reason']})" if r["reason"] else "")),
            ("Thiếu hụt vốn", f"{vnd(r['gap']) or 0} đ"),
            (
                "Số đề xuất",
                f"{vnd(r['amount_proposed']) or 0} đ, {r['tenor_weeks'] or 0} tuần, "
                f"trích {vnd(r['weekly_deduction']) or 0} đ mỗi tuần",
            ),
            ("Payout trung vị 12 tuần", f"{vnd(r['payout_base'])} đ"),
            ("PD", f"{r['pd_bps']} bps" if r["pd_bps"] is not None else "không chấm"),
            (
                "CIC",
                f"nhóm hiện tại {r['cic_group']}, xấu nhất 12 tháng {r['cic_worst_12m']}"
                if r["cic_group"] is not None
                else "không tra",
            ),
            ("Rule không đạt", ", ".join(r["failed_rules"]) or "không có"),
        ]
        if r.get("payout_warning_threshold"):
            rows.append(
                ("Ngưỡng payout cảnh báo", f"{vnd(r['payout_warning_threshold'])} đ mỗi tuần")
            )
        if r.get("statement"):
            st = r["statement"]
            rows.append(
                (
                    "Đối chiếu sao kê",
                    f"sàn {vnd(st['platform_total'])} đ, sao kê "
                    f"{vnd(st['marketplace_total'])} đ, lệch {vnd(st['difference']) or 0} đ, "
                    f"{'khớp' if st['matched'] else 'không khớp'}",
                )
            )
        return [f"**{label}**", "", "| | |", "|---|---|", *[f"| {k} | {v} |" for k, v in rows]]


def risk_model() -> RiskModel:
    try:
        import lightgbm  # noqa: F401

        return LightGbmRiskModel()
    except OSError:
        return BaselineRiskModel()


def build(d: Dossier, fonts: Fonts, expect: Expectation) -> dict[str, bytes]:
    paper, quality = d.spec.get("paper", {}), d.spec.get("scan_quality", "normal")
    unknown = set(paper) - {k.key for k in PDF_DOCUMENTS}
    if (
        unknown
        or not set(paper.values()) <= {"digital", "scan", "hand"}
        or quality not in SCAN_PROFILES
    ):
        raise SystemExit(f"{d.code}: paper không hợp lệ {paper} {quality}")
    rng = d.rng("paper")
    d.entries = LedgerBuilder(d).build()
    statement = StatementPdf(d, fonts, Typed(fonts))
    pdf = statement.render()
    d.expected = expect.run(d, pdf)
    scan = Scanned(SCAN_PROFILES["normal"])
    files = {statement.filename: pdf, PHOTO[1]: photo(pdf), SCAN[1]: scan.apply(pdf, rng)}
    docs = [
        (statement.filename, statement.label, "PDF"),
        (PHOTO[1], PHOTO[2], "ảnh chụp"),
        (SCAN[1], SCAN[2], f"PDF {scan.note}"),
    ]
    for kind in PDF_DOCUMENTS:
        mode = paper.get(kind.key, "digital")
        ink = Handwritten(fonts, d.hands[0], rng) if mode == "hand" else Typed(fonts)
        sheet = paper_for(mode, quality)
        files[kind.filename] = sheet.apply(kind(d, fonts, ink).render(), rng)
        docs.append((kind.filename, kind.label, f"PDF {sheet.note}".strip()))
    for kind in WORKBOOKS:
        files[kind.filename] = kind(d).render()
        docs.append((kind.filename, kind.label, "Excel"))
    docs.sort(key=lambda x: x[0])
    files[Readme.filename] = Readme(d, docs).render()
    files["expected.json"] = json.dumps(
        d.expected, ensure_ascii=False, indent=1, default=str
    ).encode()
    return files


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--date", default="2026-10-31")
    ap.add_argument("--out", default=str(ROOT / "eval/dossiers/out"))
    ap.add_argument("--font", default=str(BODY_FONTS / "DejaVuSans.ttf"))
    ap.add_argument("--font-bold", default=str(BODY_FONTS / "DejaVuSans-Bold.ttf"))
    ap.add_argument("--hand-fonts", default=str(ROOT / "eval/dossiers/fonts"))
    args = ap.parse_args()
    spec = json.loads((ROOT / "eval/dossiers/cases.json").read_text(encoding="utf-8"))
    sellers = {s.seller_id: s for s in SellerBook.load().sellers}
    credit = json.loads((ROOT / "api/data/cic_profiles.json").read_text(encoding="utf-8"))[
        "sellers"
    ]
    CicProfileBook.load()
    window = Window(datetime.fromisoformat(args.date).replace(hour=9, tzinfo=VN))
    fonts = Fonts.register(Path(args.font), Path(args.font_bold), Path(args.hand_fonts))
    expect = Expectation(window, risk_model())
    root = Path(args.out) / window.now.date().isoformat()
    index = [
        f"# Bộ hồ sơ mẫu Riverbank, ngày {dmy(window.now.date())}",
        "",
        f"> {spec['disclosure']}",
        "",
        "| Mã | Seller | Tình huống | Kết quả khi tạo | Sau khi nộp sao kê |",
        "|---|---|---|---|---|",
    ]
    for item in spec["cases"]:
        seller = sellers[item["seller_id"]]
        pick = random.Random(seller.seller_id).sample(fonts.hands, 2)
        d = Dossier(
            item,
            seller,
            credit.get(item["seller_id"], {}),
            window,
            lender=spec["lender"],
            account_bank=spec["account_bank"],
            disclosure=spec["disclosure"],
            hands=tuple(pick),
        )
        folder = root / d.folder
        folder.mkdir(parents=True, exist_ok=True)
        for old in folder.iterdir():
            old.unlink()
        for name, content in build(d, fonts, expect).items():
            (folder / name).write_bytes(content)
        after = d.expected.get("after_statement")
        index.append(
            f"| {d.code} | `{d.seller.seller_id}` | {item['title']} | "
            f"`{d.expected['created']['status']}` | {f'`{after["status"]}`' if after else ''} |"
        )
        print(d.folder, d.expected["created"]["status"], after and after["status"])
    (root / "README.md").write_text("\n".join(index) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
